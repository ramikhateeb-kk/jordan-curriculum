"""
teach_help.py
دليل المعلم الشامل ومساعد الشرح للحصة.
يتضمن الربط المباشر بـ Supabase للأنظمة والمحاذاة التامة RTL.
"""
import hashlib
import io
import json
import re
from urllib.parse import quote

import requests
import streamlit as st
from docx import Document
from docx.enum.table import WD_TABLE_DIRECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
from gtts import gTTS

import auth_ui
import catalog
import plan_cache
from config import GEMINI_MODEL
from rag import get_client, select_context
from word_rtl import (
    rtl_p as _rtl_p,
    rtl_table as _rtl_table,
    run as _run,
    doc_table as _doc_table,
    center as _center
)

BUCKET = "teacher-files"          # اسم الـ Bucket في Supabase Storage (يجب أن يكون Public)
TABLE = "teacher_guides"
TIMEOUT = 30
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


# ---------------------------------------------------------------- التعامل المباشر مع Supabase (REST)
def _use_supabase():
    return catalog.backend_name() == "supabase"


def get_guide_row(guide_key, lesson_title):
    """يعيد صف الدليل المحفوظ {content, file_url, file_path} أو None."""
    try:
        if _use_supabase():
            url, h = catalog._sb()
            r = requests.get(f"{url}/rest/v1/{TABLE}", headers=h, timeout=TIMEOUT, params={
                "select": "content,file_url,file_path",
                "book_slug": f"eq.{guide_key}", "lesson_title": f"eq.{lesson_title}"})
            catalog._check(r)
            rows = r.json()
            return rows[0] if rows else None
        # وضع محلي (بدون Supabase): تخزين في ملف
        content = plan_cache.get_teacher_guide(guide_key)
        return {"content": content, "file_url": None, "file_path": None} if content else None
    except Exception as e:
        st.warning(f"⚠️ تعذر الاستعلام من قاعدة البيانات: {e}")
        return None


def save_guide_row(guide_key, lesson_title, subject, content, file_url=None, file_path=None):
    """حفظ/تحديث الدليل (Upsert) مع رابط الملف إن وُجد."""
    try:
        if not _use_supabase():
            return plan_cache.save_teacher_guide(guide_key, content, subject)
        row = {"book_slug": guide_key, "lesson_title": lesson_title,
               "subject": subject, "content": content}
        if file_url:
            row["file_url"] = file_url
            row["file_path"] = file_path
        url, h = catalog._sb()
        r = requests.post(f"{url}/rest/v1/{TABLE}",
                          headers={**h, "Prefer": "resolution=merge-duplicates,return=minimal"},
                          timeout=TIMEOUT, params={"on_conflict": "book_slug,lesson_title"},
                          data=json.dumps([row]))
        catalog._check(r)
        return True
    except Exception as e:
        st.error(f"❌ خطأ أثناء الحفظ في قاعدة بيانات Supabase: {e}")
        return False


def upload_docx(path, data, download_name):
    """يرفع ملف Word إلى Supabase Storage ويعيد رابط تحميل عام مباشر."""
    url, h = catalog._sb()
    headers = {"apikey": h["apikey"], "Authorization": h["Authorization"], "Content-Type": DOCX_MIME}
    r = requests.post(f"{url}/storage/v1/object/{BUCKET}/{path}",
                      headers=headers, data=data, timeout=TIMEOUT)
    # 409 = الملف نفسه مرفوع سابقًا (المسار يتضمن بصمة المحتوى) → نعتبره نجاحًا
    if not r.ok and r.status_code != 409 and "Duplicate" not in r.text:
        raise RuntimeError(f"فشل رفع الملف ({r.status_code}): {r.text[:200]}")
    return f"{url}/storage/v1/object/public/{BUCKET}/{path}?download={quote(download_name)}"


def build_file_name(title, grade_label):
    safe = re.sub(r'[\\/:*?"<>|]', "", f"دليل_شرح_{title}_{grade_label}").replace(" ", "_")
    return safe + ".docx"


def publish_guide(guide_key, lesson_title, subject, grade_label, content):
    """ينشئ ملف Word، يرفعه إلى Storage، ويحفظ الدليل ورابطه في الجدول. يعيد الرابط أو None."""
    doc_bytes = export_to_word(lesson_title, content, grade_label, subject)
    file_link, file_path = None, None
    if _use_supabase():
        try:
            # اسم ملف ASCII فقط (Storage يرفض الحروف العربية في المسار) + بصمة المحتوى لتجنب الكاش القديم
            file_path = (f"guides/{hashlib.md5(guide_key.encode('utf-8')).hexdigest()[:16]}"
                         f"_{hashlib.md5(content.encode('utf-8')).hexdigest()[:8]}.docx")
            file_link = upload_docx(file_path, doc_bytes, build_file_name(lesson_title, grade_label))
        except Exception as e:
            st.warning(f"⚠️ تم حفظ الدليل لكن تعذر إنشاء رابط التحميل: {e}")
    save_guide_row(guide_key, lesson_title, subject, content, file_link, file_path)
    return file_link


# ---------------------------------------------------------------- معالجة العرض بتنسيق RTL
def clean_markdown_for_speech(text):
    text = re.sub(r'\*+|\#+|`+|_+', '', text)
    text = re.sub(r'\[.*?\]\(.*?\)', '', text)
    text = re.sub(r'\n+', ' ', text)
    return text.strip()


def text_to_speech(text, lang='ar'):
    try:
        clean_text = clean_markdown_for_speech(text)
        if not clean_text:
            return None
        clean_text = clean_text[:1000]
        tts = gTTS(text=clean_text, lang=lang, slow=False)
        fp = io.BytesIO()
        tts.write_to_fp(fp)
        fp.seek(0)
        return fp
    except Exception as e:
        st.error(f"خطأ في توليد الصوت: {e}")
        return None


def clean_output_text(text):
    text = re.sub(r'&lt;\s*br\s*/?&gt;', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'<\s*br\s*/?>', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'</?[^>]+>', '', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    lines = [line.strip() for line in text.split('\n')]
    return '\n'.join(lines)


def format_text_to_rtl_html(text):
    """محول النصوص الخاص بعرض المستند بتنسيق أنيق ومحاذاة يمين كاملة"""
    lines = text.split('\n')
    formatted_html = []
    in_table = False
    table_rows = []

    for line in lines:
        line_str = line.strip()

        if line_str.startswith('|'):
            in_table = True
            table_rows.append(line_str)
            continue
        else:
            if in_table:
                formatted_html.append(render_html_table(table_rows))
                in_table = False
                table_rows = []

        if not line_str or line_str == '---':
            if line_str == '---':
                formatted_html.append("<hr style='border: 1px solid #e5e7eb; margin: 15px 0;'>" )
            continue

        if line_str.startswith('### '):
            formatted_html.append(f"<h3 style='direction: rtl; text-align: right; color: #1E3A8A; margin-top: 15px;'>{line_str.replace('### ', '')}</h3>")
        elif line_str.startswith('## '):
            formatted_html.append(f"<h2 style='direction: rtl; text-align: right; color: #1E3A8A; margin-top: 18px;'>{line_str.replace('## ', '')}</h2>")
        elif line_str.startswith('# '):
            formatted_html.append(f"<h1 style='direction: rtl; text-align: right; color: #1E3A8A; margin-top: 20px;'>{line_str.replace('# ', '')}</h1>")
        elif re.match(r'^\d+\.', line_str):
            # محاذاة القوائم الرقمية
            formatted_html.append(f"<div style='direction: rtl; text-align: right; font-weight: bold; margin-top: 10px; color: #0F172A;'>{line_str}</div>")
        elif line_str.startswith('- ') or line_str.startswith('* ') or line_str.startswith('• '):
            clean_item = line_str.lstrip('-*• ').replace('**', '')
            formatted_html.append(f"<div style='direction: rtl; text-align: right; padding-right: 20px; margin-bottom: 5px; color: #334155;'>• {clean_item}</div>")
        else:
            clean_p = line_str.replace('**', '<b>').replace('**', '</b>')
            formatted_html.append(f"<p style='direction: rtl; text-align: right; margin-bottom: 8px; line-height: 1.6;'>{clean_p}</p>")

    if in_table:
        formatted_html.append(render_html_table(table_rows))

    return "\n".join(formatted_html)


def render_html_table(table_lines):
    rows = [l.strip().strip('|').split('|') for l in table_lines if '---' not in l]
    if not rows:
        return ""
    
    html = "<div style='overflow-x: auto;'><table style='direction: rtl; width: 100%; border-collapse: collapse; margin: 15px 0; border: 1px solid #cbd5e1;'>"
    for idx, row in enumerate(rows):
        tag = "th" if idx == 0 else "td"
        bg_style = "background-color: #f1f5f9; font-weight: bold; color: #1e293b;" if idx == 0 else "color: #334155;"
        html += "<tr>"
        for cell in row:
            html += f"<{tag} style='border: 1px solid #cbd5e1; padding: 10px; text-align: right; direction: rtl; {bg_style}'>{cell.strip()}</{tag}>"
        html += "</tr>"
    html += "</table></div>"
    return html


# ---------------------------------------------------------------- تصدير إلى Word
def export_to_word(title, content_text, grade_name, subject_name):
    doc = Document()
    
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    _center(doc, f"دليل شرح: {title}", 14, True)
    
    p_meta = doc.add_paragraph()
    _rtl_p(p_meta, WD_ALIGN_PARAGRAPH.RIGHT)
    _run(p_meta, f"المبحث: {subject_name}  |  الصف: {grade_name}", 10, bold=True)

    lines = content_text.split('\n')
    in_table = False
    table_lines = []

    for line in lines:
        if line.startswith('|'):
            in_table = True
            table_lines.append(line)
        else:
            if in_table:
                create_word_table(doc, table_lines)
                in_table = False
                table_lines = []

            if line.strip() and line.strip() != '---':
                p = doc.add_paragraph()
                _rtl_p(p, WD_ALIGN_PARAGRAPH.RIGHT)
                clean_line = line.replace('**', '').replace('###', '').replace('##', '').replace('#', '')
                
                if line.startswith('#'):
                    _run(p, clean_line, 12, True)
                elif line.startswith('- ') or line.startswith('* ') or line.startswith('• '):
                    _run(p, '• ' + clean_line.lstrip('-*• '), 10.5)
                else:
                    _run(p, clean_line, 10.5)

    if in_table:
        create_word_table(doc, table_lines)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def create_word_table(doc, lines):
    rows = [l.strip().strip('|').split('|') for l in lines if '---' not in l]
    if not rows or not rows[0]:
        return

    num_cols = len(rows[0])
    num_rows = len(rows)

    total_width_cm = 16.0
    col_width = total_width_cm / num_cols
    col_widths = [Cm(col_width) for _ in range(num_cols)]

    table = _doc_table(doc, num_rows, num_cols, col_widths, borders=True)
    table.table_direction = WD_TABLE_DIRECTION.RTL
    _rtl_table(table)

    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            if c_idx < len(table.columns):
                cell = table.cell(r_idx, c_idx)
                p = cell.paragraphs[0]
                _rtl_p(p, WD_ALIGN_PARAGRAPH.RIGHT)
                _run(p, val.strip(), 9.5, bold=(r_idx == 0))


# ---------------------------------------------------------------- الواجهة الرئيسية
def render_teach_help_builder(idx, slug, subject, flat_lessons, grade_label="الحادي عشر", user=None):
    """واجهة دليل الشرح: توليد → حفظ في Supabase → ملف Word برابط تحميل."""
    st.subheader(f"🧑‍🏫 مساعد الشرح (دليل المعلم الشامل) - الصف {grade_label}")

    if not flat_lessons:
        st.warning("لا توجد دروس معرفة في هذا الكتاب.")
        return

    selected_lesson_label = st.selectbox(
        "اختر الدرس لإعداد دليل مساعد الشرح له:",
        [l["label"] for l in flat_lessons],
        key="teach_help_lesson_select"
    )
    chosen_lesson = next(l for l in flat_lessons if l["label"] == selected_lesson_label)
    title = chosen_lesson["title"]

    guide_key = f"{slug}_{grade_label}_{title}".replace(" ", "_")

    # الدليل المعروض حاليًا (يبقى ظاهرًا بعد أي إعادة تشغيل للصفحة، مثل الضغط على الصوت أو التنزيل)
    shown = st.session_state.get("teach_help_shown", {})
    saved_guide = shown.get("content") if shown.get("key") == guide_key else None
    file_url = shown.get("file_url") if shown.get("key") == guide_key else None

    if st.button("📖 عرض دليل مساعد الشرح", type="primary", key="generate_teach_help_btn"):
        with st.spinner("جارٍ تحضير الدليل..."):
            # 1) أولاً: هل الدليل موجود في قاعدة البيانات؟ → استرجاعه دون أي طلب إلى Gemini
            row = get_guide_row(guide_key, title)
            if row and row.get("content"):
                saved_guide = row["content"]
                file_url = row.get("file_url")
                st.info("💡 تم استرجاع الدليل من قاعدة البيانات (لم يُولَّد من جديد).")
            else:
                # 2) غير موجود (أول طلب): توليده بـ Gemini ثم حفظه وإنشاء ملف Word ورابطه
                try:
                    pages = (chosen_lesson["start"], chosen_lesson["end"])
                    ctx_pages = select_context(idx, title, pages)
                    context_text = "\n".join([f"صفحة {p['page']}: {p['text']}" for p in ctx_pages])

                    prompt = f"""
أنت مشرف تربوي وخبير تدريس للمنهاج الأردني.
قم بإنشاء دليل شرح مبسط وموجز للمعلم لدرس: "{chosen_lesson['title']}" في مادة {subject} للصف ({grade_label}).

نص الدرس من الكتاب:
{context_text}

شروط التنسيق والصياغة والإخراج:
1. صمم المحتوى بوضوح ومباشرة وبدون تعقيد لصف ({grade_label}).
2. ممنوع استخدام وسوم HTML نهائياً.
3. التزم بعمل جداول Markdown واضحة الحقول.

المحاور المطلوبة:
1. **جدول توزيع زمن الحصة (45 دقيقة)**: جدول منظم يحتوي على الأوردة والأعمدة: (المرحلة | الزمن | الأنشطة المقترحة | الموارد والأدوات).
2. **خطوات شرح الدرس والتمايز**:
   - أولاً: المقدمة والتمهيد (5 دقائق)
   - ثانياً: العرض الشرح (20 دقيقة)
   - ثالثاً: أنشطة التمايز (ضعاف التحصيل / الطلبة المجتهدين)
3. **نموذج جدول مكتمل لأداة تقويم** (قائمة رصد Checklist أو سلم تقدير Rubric).
"""
                    client = get_client()
                    response = client.models.generate_content(model=GEMINI_MODEL, contents=prompt)
                    result_text = clean_output_text(response.text)

                    file_url = publish_guide(guide_key, title, subject, grade_label, result_text)
                    saved_guide = result_text
                    st.success("✅ تم إنشاء الدليل وحفظه في قاعدة البيانات!")
                except Exception as e:
                    st.error(f"حدث خطأ أثناء بناء الدليل: {e}")

        if saved_guide:
            st.session_state["teach_help_shown"] = {"key": guide_key, "content": saved_guide, "file_url": file_url}

    if saved_guide:
        # دليل قديم محفوظ بدون ملف → زر لإنشاء الملف والرابط (بدون طلب Gemini)
        if not file_url and _use_supabase():
            if st.button("🔗 إنشاء ملف Word ورابط تحميل لهذا الدليل", key="make_guide_link_btn"):
                with st.spinner("جارٍ إنشاء الملف ورفعه..."):
                    file_url = publish_guide(guide_key, title, subject, grade_label, saved_guide)
                    st.session_state["teach_help_shown"] = {"key": guide_key, "content": saved_guide, "file_url": file_url}

        st.markdown("---")
        st.markdown(f"### 📋 دليل مساعد الشرح للحصة - {title} ({grade_label}):")
        st.markdown(format_text_to_rtl_html(saved_guide), unsafe_allow_html=True)

        st.markdown("---")
        if file_url:
            st.link_button("🔗 تحميل الدليل (رابط مباشر)", file_url, type="primary")
            st.caption("انسخ الرابط لمشاركته:")
            st.code(file_url, language=None)

        col_download, col_audio = st.columns(2)
        with col_download:
            st.download_button(
                label="📥 تنزيل الدليل كملف Word (.docx)",
                data=export_to_word(title, saved_guide, grade_label, subject),
                file_name=build_file_name(title, grade_label),
                mime=DOCX_MIME
            )
        with col_audio:
            if st.checkbox("🔊 الاستماع لملخص الدليل"):
                with st.spinner("جارٍ توليد الصوت..."):
                    audio_fp = text_to_speech(saved_guide, lang='ar')
                    if audio_fp:
                        st.audio(audio_fp, format='audio/mp3')
