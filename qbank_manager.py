"""
qbank_manager.py
تبويب المعلم "بنك الأسئلة": يُدخل المعلم سؤالًا أو قائمة أسئلة (نص / صورة / ملف TXT أو Word)
لدرس محدد أو لكل الكتاب، فيقوم الذكاء الاصطناعي بفحصها مقابل نص الكتاب:
  1) استخراج الأسئلة وتنظيمها (طلب واحد إلى Gemini).
  2) التحقق: الارتباط بالدرس، وجود الإجابة صراحة في الكتاب، صحة الإجابة/الخيارات، رقم الصفحة المصدر
     (طلب واحد لكل 10 أسئلة). الإجابة التي كتبها المعلم لا تُرسَل لأسئلة الاختيار وصح/خطأ، بل يستنتجها
     النموذج من الكتاب ثم تُقارَن محليًا — حتى لا ينحاز النموذج لإجابة المعلم.
  3) منع التكرار محليًا (بلا Gemini) مقابل بنك الأسئلة الحالي ومقابل أسئلة الدفعة نفسها.
  4) الأسئلة السليمة وغير المكررة فقط تُضاف إلى جدول الأسئلة (qbank) فتظهر تلقائيًا في "قيمني" و"إنشاء اختبار".
في خيار "كل الكتاب" يحدد النظام الدرس المناسب لكل سؤال من رقم صفحة مصدره.
"""
import difflib
import io
import json
import re

import streamlit as st
from google.genai import types
from PIL import Image

import auth_ui
import catalog
import qbank
from config import GEMINI_MODEL
from rag import _norm, get_client, select_context

QTYPE_AR = {"mcq": "اختيار من متعدد", "tf": "صح أو خطأ", "fill": "أكمل الفراغ", "short": "سؤال قصير"}
ALL_BOOK = "كل الكتاب (يحدد النظام الدرس تلقائيًا)"

MAX_QUESTIONS = 30        # أقصى عدد أسئلة في الدفعة الواحدة
VERIFY_BATCH = 10         # عدد الأسئلة في كل طلب تحقق إلى Gemini
LESSON_CONTEXT_BUDGET = 40000
BOOK_CTX_BUDGET = 70000   # نص الكتاب الأقصى المُرسَل في طلب تحقق واحد عند اختيار "كل الكتاب"
PER_Q_CHARS, PER_Q_PAGES = 9000, 4
MAX_FILE_BYTES = 5 * 1024 * 1024
DUP_RATIO = 0.92

PARSE_SYSTEM = (
    "أنت مساعد يستخرج الأسئلة من نص أو صورة ويحوّلها إلى JSON منظَّم. التزم بالآتي بصرامة:\n"
    "1) انقل صياغة كل سؤال كما هي دون تغيير، ولا تضف أي سؤال أو خيار أو إجابة من عندك.\n"
    "2) حدّد نوع كل سؤال: mcq (له خيارات أ/ب/ج/د)، tf (عبارة يُحكم عليها بصح أو خطأ)، "
    "fill (جملة فيها فراغ؛ اكتب الفراغ دائمًا ______)، short (سؤال يُجاب بجملة).\n"
    "3) اكتب الخيارات بلا حروف الترقيم (أ، ب، ...).\n"
    "4) لا تضع الإجابة الصحيحة إلا إذا كانت مُعلَّمة صراحةً في المدخل (مفتاح إجابات، علامة ✓، تسطير، "
    "«الإجابة:»)؛ وإلا اجعلها null.\n"
    "5) تجاهل أي نص ليس سؤالًا (عناوين، تعليمات عامة). أعد JSON فقط بلا أي نص آخر."
)
PARSE_SCHEMA = (
    '[{"type":"mcq","question":"...","options":["...","...","...","..."],"correct_index":null,'
    '"correct":null,"model_answer":null},'
    '{"type":"tf","question":"عبارة ...","options":null,"correct_index":null,"correct":null,"model_answer":null},'
    '{"type":"fill","question":"... ______ ...","options":null,"correct_index":null,"correct":null,'
    '"model_answer":null},'
    '{"type":"short","question":"...","options":null,"correct_index":null,"correct":null,"model_answer":null}]'
)

VERIFY_SYSTEM = (
    "أنت مدقّق أسئلة صارم للمنهاج الأردني. مصدرك الوحيد نص الكتاب المرفق، ولا تستخدم أي معلومة من خارجه "
    "حتى لو كانت صحيحة في الواقع، ولا تخمّن. لكل سؤال حدّد:\n"
    "- related: هل يتعلق السؤال بمحتوى النص المرفق/النطاق المطلوب؟\n"
    "- supported: هل إجابة السؤال موجودة صراحة في النص (لا مستنتجة من خارجه)؟\n"
    "- page: رقم الصفحة التي تؤكد الإجابة، من أرقام الصفحات الموجودة فعلًا في النص المرفق.\n"
    "- للاختيار من متعدد: book_correct_index (رقم الخيار الصحيح بدءًا من 0 حسب الكتاب)، و"
    "other_options_wrong (true إن كانت الخيارات الثلاثة الأخرى خاطئة أو غير مدعومة بالكتاب، "
    "وfalse إن كان أكثر من خيار صحيحًا).\n"
    "- لصح/خطأ: book_correct (true/false حسب الكتاب).\n"
    "- لأكمل الفراغ والسؤال القصير: book_answer (الإجابة حسب الكتاب بإيجاز)، وإن وُجدت teacher_answer "
    "فحدّد teacher_answer_correct (true إن كانت صحيحة المعنى حسب الكتاب، ويُقبل اختلاف الصياغة).\n"
    "- explanation: توضيح قصير من الكتاب، وreason: سبب موجز بالعربية إن كان السؤال غير مرتبط أو غير مدعوم.\n"
    "أعد JSON فقط بلا أي نص آخر."
)
VERIFY_SCHEMA = (
    '[{"i":0,"related":true,"supported":true,"page":12,"book_correct_index":0,"other_options_wrong":true,'
    '"book_correct":null,"book_answer":null,"teacher_answer_correct":null,"explanation":"...","reason":""}]'
)


# ---------------------------------------------------------------- أدوات عامة
def _ask_json(client, system, contents):
    for _ in range(2):
        try:
            r = client.models.generate_content(
                model=GEMINI_MODEL, contents=contents,
                config=types.GenerateContentConfig(
                    system_instruction=system, temperature=0,
                    response_mime_type="application/json"))
            txt = re.sub(r"^```(?:json)?|```$", "", r.text.strip()).strip()
            return json.loads(txt)
        except Exception as e:  # noqa: BLE001
            print("خطأ في توليد/قراءة JSON:", e)
    return None


def read_upload(f):
    """يعيد (نص, صورة PIL أو None) من ملف مرفوع: TXT / DOCX / صورة."""
    name = (f.name or "").lower()
    if f.size > MAX_FILE_BYTES:
        raise ValueError("حجم الملف كبير جدًا. الحد الأقصى المسموح به هو 5MB.")
    if name.endswith(".txt"):
        raw = f.read()
        for enc in ("utf-8-sig", "utf-8", "cp1256"):
            try:
                return raw.decode(enc), None
            except UnicodeDecodeError:
                continue
        raise ValueError("تعذّر قراءة ترميز الملف النصي.")
    if name.endswith(".docx"):
        from docx import Document
        d = Document(io.BytesIO(f.read()))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts), None
    return "", Image.open(f)


def _norm_q(t):
    t = _norm(str(t))
    t = re.sub(r"[^\w\s]|_", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def _lesson_for_page(flat_lessons, page):
    return next((l for l in flat_lessons if l["start"] <= page <= l["end"]), None)


# ---------------------------------------------------------------- 1) الاستخراج + فحص الشكل محليًا
def parse_questions(client, raw_text, image):
    prompt = (f"{PARSE_SYSTEM}\n\nاستخرج كل الأسئلة من المدخل التالي بهذا الشكل بالضبط:\n{PARSE_SCHEMA}\n\n"
              f"المدخل:\n{raw_text or '(الأسئلة في الصورة المرفقة)'}")
    contents = [image, prompt] if image is not None else prompt
    data = _ask_json(client, PARSE_SYSTEM, contents)
    return data if isinstance(data, list) else None


def clean_question(item):
    """يفحص بنية السؤال محليًا. يعيد (سؤال منظَّم أو None، سبب الرفض)."""
    if not isinstance(item, dict):
        return None, "تعذّر فهم هذا السؤال."
    qtype = item.get("type")
    text = str(item.get("question") or "").strip()
    if qtype not in QTYPE_AR:
        return None, "نوع السؤال غير معروف."
    if not text:
        return None, "نص السؤال فارغ."
    q = {"type": qtype, "question": text}
    if qtype == "mcq":
        opts = item.get("options")
        if not (isinstance(opts, list) and len(opts) == 4 and all(str(o).strip() for o in opts)
                and len({_norm_q(o) for o in opts}) == 4):
            return None, "سؤال الاختيار من متعدد يجب أن يحتوي على 4 خيارات مختلفة (أ، ب، ج، د)."
        q["options"] = [str(o).strip() for o in opts]
        ci = item.get("correct_index")
        q["correct_index"] = ci if isinstance(ci, int) and not isinstance(ci, bool) and 0 <= ci < 4 else None
    elif qtype == "tf":
        c = item.get("correct")
        q["correct"] = c if isinstance(c, bool) else None
    else:
        if qtype == "fill":
            q["question"] = re.sub(r"_{2,}|\.{4,}|…+", "______", text)
            if "______" not in q["question"]:
                return None, "سؤال أكمل الفراغ يجب أن يحتوي على فراغ (______)."
        ma = str(item.get("model_answer") or "").strip()
        q["model_answer"] = ma or None
    return q, ""


# ---------------------------------------------------------------- 2) التحقق من الكتاب
def _context_for(idx, questions, lesson):
    if lesson:
        query = " ".join(q["question"] for q in questions)
        return select_context(idx, query, (lesson["start"], lesson["end"]), max_chars=LESSON_CONTEXT_BUDGET)
    seen, used = {}, 0
    for q in questions:
        query = q["question"] + " " + " ".join(q.get("options") or [])
        for p in select_context(idx, query, None, max_chars=PER_Q_CHARS, k=PER_Q_PAGES):
            if p["page"] not in seen and used + len(p["text"]) <= BOOK_CTX_BUDGET:
                seen[p["page"]] = p
                used += len(p["text"])
    return sorted(seen.values(), key=lambda p: p["page"])


def verify_batch(client, idx, batch, lesson):
    """يتحقق من دفعة أسئلة (طلب واحد). يعيد (قائمة نتائج حسب الترتيب، صفحات النص المُرسَل)."""
    ctx_pages = _context_for(idx, batch, lesson)
    if not ctx_pages:
        return [None] * len(batch), set()
    ctx = "\n\n".join(f"=== صفحة {p['page']} ===\n{p['text']}" for p in ctx_pages)
    items = []
    for i, q in enumerate(batch):
        it = {"i": i, "type": q["type"], "question": q["question"]}
        if q["type"] == "mcq":
            it["options"] = q["options"]          # لا نرسل إجابة المعلم لأسئلة الاختيار وصح/خطأ
        if q["type"] in ("fill", "short") and q.get("model_answer"):
            it["teacher_answer"] = q["model_answer"]
        items.append(it)
    scope = (f"النطاق: الدرس «{lesson['title']}» (الصفحات {lesson['start']}–{lesson['end']})؛ "
             "السؤال المرتبط هو ما يخص هذا الدرس فقط." if lesson else
             "النطاق: كامل الكتاب؛ السؤال المرتبط هو ما يخص محتوى الكتاب.")
    prompt = (f"{scope}\n\nنص الكتاب (المصدر الوحيد):\n{ctx}\n\n"
              f"الأسئلة المطلوب فحصها:\n{json.dumps(items, ensure_ascii=False)}\n\n"
              f"أعد JSON (عنصر لكل سؤال بنفس رقم i) بهذا الشكل:\n{VERIFY_SCHEMA}")
    data = _ask_json(client, VERIFY_SYSTEM, prompt)
    by_i = {}
    if isinstance(data, list):
        for v in data:
            if isinstance(v, dict) and isinstance(v.get("i"), int):
                by_i[v["i"]] = v
    return [by_i.get(i) for i in range(len(batch))], {p["page"] for p in ctx_pages}


def decide(q, v, valid_pages):
    """يقرّر قبول السؤال محليًا بناءً على نتيجة التحقق. يعيد (قُبل؟، ملاحظة/سبب، السؤال النهائي)."""
    if not isinstance(v, dict):
        return False, "تعذّر فحص هذا السؤال (لم يصل رد صالح من الذكاء الاصطناعي). أعد المحاولة.", None
    if v.get("related") is not True:
        return False, v.get("reason") or "السؤال غير مرتبط بالدرس/الكتاب المحدد.", None
    if v.get("supported") is not True:
        return False, v.get("reason") or "إجابة السؤال غير موجودة صراحة في نص الكتاب.", None
    page = v.get("page")
    if not (isinstance(page, int) and not isinstance(page, bool) and page in valid_pages):
        return False, "تعذّر تحديد صفحة مصدر صحيحة للسؤال في الكتاب.", None

    final, note = dict(q), ""
    if q["type"] == "mcq":
        bi = v.get("book_correct_index")
        if not (isinstance(bi, int) and not isinstance(bi, bool) and 0 <= bi < 4):
            return False, "تعذّر تحديد الخيار الصحيح من الكتاب.", None
        if v.get("other_options_wrong") is not True:
            return False, "قد يوجد أكثر من خيار صحيح حسب الكتاب؛ عدّل الخيارات ثم أعد المحاولة.", None
        if q.get("correct_index") is not None and q["correct_index"] != bi:
            return False, (f"الإجابة المحددة («{q['options'][q['correct_index']]}») تخالف الكتاب؛ "
                           f"الإجابة الصحيحة حسب الكتاب: «{q['options'][bi]}» (صفحة {page})."), None
        if q.get("correct_index") is None:
            final["correct_index"], note = bi, "استُنتجت الإجابة الصحيحة من الكتاب"
    elif q["type"] == "tf":
        bc = v.get("book_correct")
        if not isinstance(bc, bool):
            return False, "تعذّر الحكم على العبارة من الكتاب.", None
        if q.get("correct") is not None and q["correct"] != bc:
            return False, (f"الحكم المحدد ({'صح' if q['correct'] else 'خطأ'}) يخالف الكتاب؛ "
                           f"الحكم الصحيح: {'صح' if bc else 'خطأ'} (صفحة {page})."), None
        if q.get("correct") is None:
            final["correct"], note = bc, "استُنتج الحكم الصحيح من الكتاب"
    else:
        ba = str(v.get("book_answer") or "").strip()
        if q.get("model_answer"):
            if v.get("teacher_answer_correct") is not True:
                return False, (f"الإجابة المكتوبة غير مطابقة للكتاب"
                               + (f"؛ الإجابة حسب الكتاب: «{ba}»" if ba else "") + f" (صفحة {page})."), None
        elif ba:
            final["model_answer"], note = ba, "استُنتجت الإجابة النموذجية من الكتاب"
        else:
            return False, "تعذّر تحديد الإجابة النموذجية من الكتاب.", None
    final["explanation"] = str(v.get("explanation") or "").strip()
    final["page"] = page
    final["source"] = "teacher"
    return True, note, final


# ---------------------------------------------------------------- 3) منع التكرار
class DupIndex:
    """يقارن نص السؤال (بعد التطبيع) مع أسئلة البنك لنفس الدرس، دون أي طلب إلى Gemini."""

    def __init__(self, slug):
        self.slug, self.cache = slug, {}

    def _texts(self, lesson):
        key = (lesson["start"], lesson["end"])
        if key not in self.cache:
            texts = []
            for qt in QTYPE_AR:
                texts += [_norm_q(d.get("question", "")) for _, d in
                          qbank.list_questions(self.slug, key[0], key[1], qt)]
            self.cache[key] = texts
        return self.cache[key]

    def is_dup(self, lesson, text):
        n = _norm_q(text)
        return any(e == n or difflib.SequenceMatcher(None, e, n).ratio() >= DUP_RATIO
                   for e in self._texts(lesson))

    def remember(self, lesson, text):
        self._texts(lesson).append(_norm_q(text))


# ---------------------------------------------------------------- خط المعالجة الكامل
def process(client, idx, slug, flat_lessons, scope_lesson, raw_text, image, user=None):
    """يعيد (قائمة نتائج، ملاحظات عامة). كل نتيجة: n, question, type, status, reason, lesson, page."""
    results, notes = [], []
    if user is not None:
        auth_ui.require_quota_or_stop(user)
    parsed = parse_questions(client, raw_text, image)
    if not parsed:
        return [], ["تعذّر استخراج أي سؤال من المدخل. تأكد من وضوح النص/الصورة ثم أعد المحاولة."]
    if len(parsed) > MAX_QUESTIONS:
        notes.append(f"تم فحص أول {MAX_QUESTIONS} سؤالًا فقط من أصل {len(parsed)}؛ أرسل الباقي في دفعة أخرى.")
        parsed = parsed[:MAX_QUESTIONS]

    dups = DupIndex(slug)
    pending = []   # (رقم السؤال، السؤال المنظَّم)
    for n, item in enumerate(parsed, start=1):
        q, err = clean_question(item)
        label = str((item or {}).get("question", "") if isinstance(item, dict) else item)[:200]
        if not q:
            results.append({"n": n, "question": label, "type": (item or {}).get("type") if isinstance(item, dict) else None,
                            "status": "rejected", "reason": err, "lesson": None, "page": None})
            continue
        if scope_lesson and dups.is_dup(scope_lesson, q["question"]):   # نوفّر طلب التحقق للمكرر
            results.append({"n": n, "question": q["question"], "type": q["type"], "status": "duplicate",
                            "reason": "السؤال موجود مسبقًا في بنك أسئلة هذا الدرس.",
                            "lesson": scope_lesson["title"], "page": None})
            continue
        pending.append((n, q))

    for s in range(0, len(pending), VERIFY_BATCH):
        chunk = pending[s:s + VERIFY_BATCH]
        if user is not None:
            auth_ui.require_quota_or_stop(user)
        verdicts, valid_pages = verify_batch(client, idx, [q for _, q in chunk], scope_lesson)
        for (n, q), v in zip(chunk, verdicts):
            base = {"n": n, "question": q["question"], "type": q["type"], "lesson": None, "page": None}
            ok, info, final = decide(q, v, valid_pages)
            if not ok:
                results.append({**base, "status": "rejected", "reason": info})
                continue
            lesson = scope_lesson or _lesson_for_page(flat_lessons, final["page"])
            if lesson is None:
                results.append({**base, "status": "rejected", "page": final["page"],
                                "reason": f"صفحة المصدر ({final['page']}) لا تقع ضمن أي درس معرّف في هذا الكتاب."})
                continue
            base.update(lesson=lesson["title"], page=final["page"])
            if dups.is_dup(lesson, q["question"]):
                results.append({**base, "status": "duplicate",
                                "reason": "السؤال موجود مسبقًا في بنك أسئلة هذا الدرس."})
                continue
            try:
                qbank.add_question(slug, lesson["start"], lesson["end"], q["type"], final)
                dups.remember(lesson, q["question"])
                results.append({**base, "status": "added", "reason": info})
            except Exception as e:  # noqa: BLE001
                results.append({**base, "status": "error", "reason": f"تعذّر الحفظ في قاعدة البيانات: {e}"})
    results.sort(key=lambda r: r["n"])
    return results, notes


# ---------------------------------------------------------------- استعراض/تعديل/حذف أسئلة المبحث (مشرف/مدير)
def render_manage_existing(subject, flat_lessons):
    st.markdown("### 📂 استعراض أسئلة هذا المبحث وتعديلها أو حذفها")
    try:
        books = [b for b in catalog.list_books() if b["subject"] == subject]
    except Exception as e:  # noqa: BLE001
        st.warning(f"تعذّر تحميل كتب هذا المبحث: {e}")
        return
    slug_grade = {b["slug"]: b["grade"] for b in books}
    rows = qbank.list_by_slugs(list(slug_grade))
    if not rows:
        st.caption("لا توجد أسئلة محفوظة لهذا المبحث بعد.")
        return
    rows.sort(key=lambda r: (slug_grade.get(r[1], 0), r[2]))
    for qid, slug_, ps, pe, qtype, data in rows:
        grade = slug_grade.get(slug_, "؟")
        with st.expander(f"الصف {grade} (ص {ps}–{pe}) — {QTYPE_AR.get(qtype, qtype)} — "
                         f"{data.get('question', '')[:70]}"):
            new_text = st.text_area("نص السؤال", value=data.get("question", ""), key=f"qm_text_{qid}")
            new_data = dict(data, question=new_text)
            if qtype == "mcq":
                opts = data.get("options") or ["", "", "", ""]
                new_opts = [st.text_input(f"الخيار {i+1}", value=o, key=f"qm_opt_{qid}_{i}")
                           for i, o in enumerate(opts)]
                ci = st.number_input("رقم الخيار الصحيح (0 للأول)", min_value=0, max_value=3,
                                     value=int(data.get("correct_index") or 0), key=f"qm_ci_{qid}")
                new_data.update(options=new_opts, correct_index=int(ci))
            elif qtype == "tf":
                c = st.radio("الحكم الصحيح", ["صح", "خطأ"],
                            index=0 if data.get("correct") else 1, key=f"qm_tf_{qid}")
                new_data["correct"] = c == "صح"
            elif qtype in ("fill", "short"):
                new_data["model_answer"] = st.text_input("الإجابة النموذجية", value=data.get("model_answer") or "",
                                                          key=f"qm_ans_{qid}")
            c1, c2 = st.columns(2)
            if c1.button("💾 حفظ التعديل", key=f"qm_save_{qid}"):
                qbank.update_question(qid, new_data)
                st.success("تم حفظ التعديل.")
                st.rerun()
            if c2.button("🗑️ حذف السؤال", key=f"qm_del_{qid}"):
                qbank.delete_question(qid)
                st.success("تم حذف السؤال.")
                st.rerun()


# ---------------------------------------------------------------- الواجهة
STATUS_ICON = {"added": "✅", "duplicate": "♻️", "rejected": "❌", "error": "⚠️"}


def render_qbank_builder(idx, slug, subject, flat_lessons, subject_scope=None, user=None):
    """subject_scope: None = بلا قيد (المدير العام)، مجموعة مباحث = مشرف مقيَّد بها، مجموعة فارغة = ممنوع."""
    st.subheader("🏦 بنك الأسئلة")
    if subject_scope is not None and subject not in subject_scope:
        st.error("لا تملك صلاحية الإشراف على هذا المبحث.")
        return

    st.caption("أدخل سؤالًا أو قائمة أسئلة (نص أو صورة أو ملف). يفحصها الذكاء الاصطناعي مقابل الكتاب، "
               "وتُضاف إلى البنك الأسئلة السليمة وغير المكررة فقط.")

    render_manage_existing(subject, flat_lessons)
    st.markdown("---")
    st.markdown("### ➕ إضافة أسئلة جديدة")

    if not flat_lessons:
        st.info("لم تُدخل دروس هذا الكتاب بعد.")
        return

    scope_label = st.selectbox("نطاق الأسئلة", [ALL_BOOK] + [l["label"] for l in flat_lessons],
                               key="qbm_scope")
    scope_lesson = None if scope_label == ALL_BOOK else next(l for l in flat_lessons if l["label"] == scope_label)

    raw_text = st.text_area(
        "اكتب الأسئلة هنا (يمكن كتابة عدة أسئلة من أنواع مختلفة، ويُفضَّل ذكر الإجابة الصحيحة إن وُجدت)",
        height=200, key="qbm_text",
        placeholder="مثال:\n1) تأسست إمارة شرق الأردن عام: أ) 1921  ب) 1946  ج) 1950  د) 1923\n"
                    "2) (صح/خطأ) عاصمة المملكة هي عمّان.\n3) أكمل: استقلت المملكة عام ______")
    up = st.file_uploader("أو ارفع صورة (PNG/JPG) أو ملف نصي (TXT) أو ملف Word (DOCX) يحتوي الأسئلة",
                          type=["png", "jpg", "jpeg", "txt", "docx"], key="qbm_file")
    if up is not None and up.type and up.type.startswith("image/"):
        st.image(up, width=260)

    if st.button("🔍 افحص وأضف إلى البنك", type="primary", key="qbm_go"):
        if not raw_text.strip() and up is None:
            st.warning("يرجى كتابة الأسئلة أو رفع ملف/صورة أولًا.")
        else:
            with st.spinner("جارٍ استخراج الأسئلة وفحصها مقابل الكتاب..."):
                try:
                    f_text, image = ("", None)
                    if up is not None:
                        f_text, image = read_upload(up)
                    full_text = "\n".join(t for t in (raw_text.strip(), f_text.strip()) if t)
                    client = get_client()   # نحتفظ بالمتغير كي لا يُغلق الاتصال أثناء الطلبات
                    results, notes = process(client, idx, slug, flat_lessons, scope_lesson, full_text, image, user=user)
                    st.session_state["qbm_results"] = {"results": results, "notes": notes}
                except Exception as e:  # noqa: BLE001
                    st.session_state.pop("qbm_results", None)
                    st.error(f"حدث خطأ أثناء المعالجة: {e}")

    out = st.session_state.get("qbm_results")
    if not out:
        return
    for n in out["notes"]:
        st.info(n)
    res = out["results"]
    if res:
        counts = {k: sum(r["status"] == k for r in res) for k in STATUS_ICON}
        (st.success if counts["added"] else st.warning)(
            f"✅ أُضيف: {counts['added']}   |   ♻️ مكرر: {counts['duplicate']}   |   "
            f"❌ مرفوض: {counts['rejected']}" + (f"   |   ⚠️ خطأ: {counts['error']}" if counts["error"] else ""))
        for r in res:
            where = f" — «{r['lesson']}» (ص {r['page']})" if r.get("lesson") and r.get("page") else \
                    (f" — «{r['lesson']}»" if r.get("lesson") else "")
            with st.expander(f"{STATUS_ICON[r['status']]} {r['n']}. {r['question'][:80]}{where}",
                             expanded=r["status"] != "added"):
                if r.get("type"):
                    st.caption(f"النوع: {QTYPE_AR.get(r['type'], r['type'])}")
                st.write(r["reason"] or ("تمت إضافة السؤال إلى البنك." if r["status"] == "added" else ""))
    if st.button("🧹 إخفاء النتائج", key="qbm_clear"):
        st.session_state.pop("qbm_results", None)
        st.rerun()
