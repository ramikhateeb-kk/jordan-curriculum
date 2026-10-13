"""
ميزة تحليل المحتوى المطابقة للملف المرفق (FORM # QF71-1-47rev)
"""
import io
import json
import re

import streamlit as st
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from google.genai import types

import accounts as acc
import auth_ui
import english_docs as en
import plan_cache
import profile_ctx
from config import GEMINI_MODEL
from rag import get_client, select_context
from word_rtl import center as _center, doc_table as _doc_table
from word_rtl import rtl_p as _rtl_p, rtl_section as _rtl_section
from word_rtl import rtl_table as _rtl_table, run as _run

SYSTEM_ANALYSIS = (
    "أنت مشرف تربوي خبير يحلل محتوى الكتب المدرسية الأردنية. قم بتقديم تحليل محتوى دقيق وشامل للوحدة "
    "يشمل: النتاجات، المفاهيم والمصطلحات، التعميمات، الرموز والأشكال، والمهارات. أعد JSON فقط."
)

SCHEMA_ANALYSIS = '''{
  "outcomes": ["توضيح مفهوم...", "وصف مكونات...", "تحديد أبرز التحديات..."],
  "concepts": ["الحوسبة السحابية", "التخزين السحابي", "أخلاقيات السلامة"],
  "generalizations": ["التعامل الصحيح مع شبكة الإنترنت", "مراعاة السلامة العامة في المختبر"],
  "symbols_shapes": [],
  "skills": ["تطبق مفهوم الحوسبة السحابية عمليا في حياتنا"]
}'''


def _generate_unit_analysis(idx, unit, lang="ar"):
    client = get_client()
    ctx_pages = select_context(idx, unit["title"], (unit["start"], unit["end"]), max_chars=35000)
    if not ctx_pages:
        return None
    if lang == "en":
        ctx = "\n\n".join(f"=== Page {p['page']} ===\n{p['text']}" for p in ctx_pages)
        system = en.SYSTEM_ANALYSIS
        prompt = (f"Unit: {unit['title']}\n\nFull unit text:\n{ctx}\n\n"
                  f"Analyse the content of this unit following this schema:\n{en.SCHEMA_ANALYSIS}")
    else:
        ctx = "\n\n".join(f"=== صفحة {p['page']} ===\n{p['text']}" for p in ctx_pages)
        system = SYSTEM_ANALYSIS
        prompt = f"الوحدة: {unit['title']}\n\nنص الوحدة كاملاً:\n{ctx}\n\nقم بتحليل محتوى هذه الوحدة طبقاً للسكيما التالية:\n{SCHEMA_ANALYSIS}"
    try:
        r = client.models.generate_content(
            model=GEMINI_MODEL, contents=prompt,
            config=types.GenerateContentConfig(system_instruction=system, temperature=0.3, response_mime_type="application/json"))
        return json.loads(re.sub(r"^```(?:json)?|```$", "", r.text.strip()).strip())
    except Exception as e:
        print("خطأ أثناء تحليل المحتوى:", e)
        return None


def get_or_generate_content_analysis(idx, slug, units, user=None, lang="ar"):
    cached = plan_cache.get_content_analysis(slug)
    if cached:
        return cached, True
    if user is not None:
        auth_ui.require_quota_or_stop(user)   # خصم واحد لكل عملية توليد

    units_analysis = []
    for unit in units:
        analysis_data = _generate_unit_analysis(idx, unit, lang)
        if analysis_data:
            units_analysis.append({
                "title": unit["title"],
                "pages": f"{unit.get('start', '')} - {unit.get('end', '')}",
                "analysis": analysis_data
            })

    if units_analysis:
        try:
            plan_cache.save_content_analysis(slug, units_analysis)
        except Exception as e:
            print("تعذر حفظ تحليل المحتوى:", e)

    return units_analysis, False


def _add_analysis_page(doc, first_page, meta, unit_data):
    section = doc.sections[0] if first_page else doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    if section.page_width < section.page_height:
        section.page_width, section.page_height = section.page_height, section.page_width

    section.top_margin = section.bottom_margin = Cm(0.8)
    section.left_margin = section.right_margin = Cm(2.0)
    _rtl_section(section)

    # التذييل
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.paragraph_format.space_before = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    _run(_rtl_p(footer_p, WD_ALIGN_PARAGRAPH.RIGHT), "FROM # QF71-1-47rev", 8, False)

    _center(doc, f"تحليل محتوى خطة الفصل الدراسي {meta['semester']} لمادة {meta['subject']}", 15, True)

    top_table = _doc_table(doc, 2, 2, [Cm(12.5), Cm(12.5)])
    _rtl_table(top_table)

    r0 = top_table.rows[0].cells
    _run(_rtl_p(r0[0].paragraphs[0], WD_ALIGN_PARAGRAPH.RIGHT), f"الصف : {meta['grade']}", 13, True)
    _run(_rtl_p(r0[1].paragraphs[0], WD_ALIGN_PARAGRAPH.LEFT), f"للعام الدراسي {meta['year']}", 13, True)

    r1 = top_table.rows[1].cells
    _run(_rtl_p(r1[0].paragraphs[0], WD_ALIGN_PARAGRAPH.RIGHT), f"عنوان الوحدة : {unit_data['title']}", 13, True)
    _run(_rtl_p(r1[1].paragraphs[0], WD_ALIGN_PARAGRAPH.LEFT), f"الصفحات : {unit_data.get('pages', '')}", 13, True)

    # الجدول (6 أعمدة)
    table = _doc_table(doc, 2, 6, [Cm(3.0), Cm(7.0), Cm(5.0), Cm(4.0), Cm(3.0), Cm(3.0)], borders=True)
    _rtl_table(table)

    headers = ["الوحدة", "النتاجات", "المفاهيم و المصطلحات", "التعميمات", "الرموز والأشكال", "المهارات"]
    for i, h in enumerate(headers):
        _run(_rtl_p(table.rows[0].cells[i].paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER), h, 13, True)

    analysis = unit_data["analysis"]
    row_cells = table.rows[1].cells

    _run(_rtl_p(row_cells[0].paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER), unit_data['title'], 13, True)

    p_out = row_cells[1].paragraphs[0]
    _run(_rtl_p(p_out, WD_ALIGN_PARAGRAPH.RIGHT), "يتوقع في نهاية الوحدة أن يكون الطالب قادراً على :", 13, True)
    for item in analysis.get("outcomes", []):
        _run(_rtl_p(row_cells[1].add_paragraph(), WD_ALIGN_PARAGRAPH.RIGHT), f"- {item}", 13)

    for item in analysis.get("concepts", []):
        _run(_rtl_p(row_cells[2].add_paragraph(), WD_ALIGN_PARAGRAPH.RIGHT), f"{item}", 13)

    for item in analysis.get("generalizations", []):
        _run(_rtl_p(row_cells[3].add_paragraph(), WD_ALIGN_PARAGRAPH.RIGHT), f"- {item}", 13)

    for item in analysis.get("symbols_shapes", []):
        _run(_rtl_p(row_cells[4].add_paragraph(), WD_ALIGN_PARAGRAPH.RIGHT), f"- {item}", 13)

    for item in analysis.get("skills", []):
        _run(_rtl_p(row_cells[5].add_paragraph(), WD_ALIGN_PARAGRAPH.RIGHT), f"- {item}", 13)


def build_analysis_docx(meta, units_analysis, lang="ar"):
    if lang == "en":
        return en.build_analysis_docx(meta, units_analysis)
    doc = Document()
    for i, u_data in enumerate(units_analysis):
        _add_analysis_page(doc, i == 0, meta, u_data)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_content_analysis_builder(idx, slug, subject, units, user=None, ctx=None, lang="ar"):
    st.subheader("📊 تحليل محتوى المنهاج" if lang != "en" else "📊 تحليل المحتوى (Content Analysis — English, LTR)")
    ctx = ctx or {}
    if lang == "en":
        st.caption("🇬🇧 هذا كتاب لغة إنجليزية: سيُنتَج التحليل بالإنجليزية واتجاه الصفحة من اليسار إلى اليمين.")
    if plan_cache.get_content_analysis(slug) is not None:
        st.info("💡 يوجد تحليل محتوى محفوظ لهذا المبحث وسيتم استرجاعه فوراً دون طلب جديد لـ Gemini.")

    grade_txt = (f"Grade {ctx['grade']}" + (f" ({ctx['track']})" if ctx.get("track") else "")) \
        if lang == "en" else ctx.get("class_label", "")
    grade = profile_ctx.known_input("الصف الدراسي", grade_txt, "ca_grade")
    year = profile_ctx.known_input("العام الدراسي", ctx.get("greg_year"), "ca_year", placeholder="2026/2027")
    profile_ctx.known_summary([("المبحث", subject), ("الصف", grade_txt), ("الفصل", ctx.get("sem_name")),
                               ("العام الدراسي", ctx.get("greg_year"))])

    if st.button("🚀 إنشاء / استرجاع تحليل المحتوى", type="primary", key="btn_ca"):
        with st.spinner("جارٍ معالجة تحليل المحتوى..."):
            units_analysis, from_cache = get_or_generate_content_analysis(idx, slug, units, user=user, lang=lang)

        if units_analysis:
            semester = ctx.get("sem_name_en") if lang == "en" else ctx.get("sem_name", "")
            meta = {"subject": subject, "grade": grade, "year": year, "semester": semester}
            docx_bytes = build_analysis_docx(meta, units_analysis, lang)
            st.success("تم إعداد تحليل المحتوى بنجاح!")
            st.download_button(
                "⬇️ تنزيل تحليل المحتوى (Word)",
                data=docx_bytes,
                file_name=f"تحليل_محتوى_{subject}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
