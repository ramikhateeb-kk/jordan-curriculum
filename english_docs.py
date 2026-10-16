"""
english_docs.py
نسخة إنجليزية (من اليسار إلى اليمين LTR) لثلاثة نماذج: خطة الدرس، الخطة الفصلية، تحليل المحتوى — لكتب
اللغة الإنجليزية (book["lang"] == "en"). تعكس ترتيب أعمدة النماذج العربية: ما كان أقصى اليمين صار أقصى اليسار.
تحتوي أيضًا تعليمات الذكاء الاصطناعي ومخططات JSON والبيانات الاحتياطية بالإنجليزية.
"""
import io
import re

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


FONT = "Arial"
L, C, R = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT

# ---------------------------------------------------------------- تعليمات الذكاء الاصطناعي
SYSTEM_LESSON = (
    "You are an expert educational supervisor helping teachers of the Jordanian English curriculum write lesson "
    "plans. Base the content ONLY on the attached lesson text and do not go beyond its topic. Write clear, practical "
    "English, concise (each cell must not exceed 2-3 lines), using action verbs such as: presents, asks, discusses... "
    "Do not be wordy so that the plan fits on one page. Return JSON only, no other text."
)
SCHEMA_LESSON = '''{
  "prior_learning": "one focused sentence about prior learning",
  "outcomes": ["focused outcome 1", "focused outcome 2", "focused outcome 3"],
  "stages": [
    {"teacher_role": "teacher's role (2-3 lines)", "student_role": "student's role (2-3 lines)"},
    {"teacher_role": "teacher's role (2-3 lines)", "student_role": "student's role (2-3 lines)"},
    {"teacher_role": "teacher's role (2-3 lines)", "student_role": "student's role (2-3 lines)"},
    {"teacher_role": "teacher's role (2-3 lines)", "student_role": "student's role (2-3 lines)"}
  ]
}'''
SYSTEM_SEMESTER = (
    "You are an expert educational supervisor helping teachers of the Jordanian English curriculum write the "
    "semester plan. Use ONLY the attached unit text. Write short, precise bullet-style items in English so that the "
    "unit fits on a single page. Return JSON only."
)
SCHEMA_SEMESTER = '''{
  "outcomes": ["Identify the key vocabulary of the unit", "Use the target grammar in context"],
  "resources": ["Student's book", "Computer lab", "Data show"],
  "strategies": ["Direct instruction", "Cooperative learning", "Role play"],
  "eval_strategies": ["Direct observation", "Performance-based assessment"],
  "eval_tools": ["Rating scale", "Checklist"],
  "activities": ["Worksheets", "Interactive presentations"]
}'''
SYSTEM_ANALYSIS = (
    "You are an expert educational supervisor analysing the content of Jordanian school English textbooks. Give an "
    "accurate, complete content analysis of the unit covering: outcomes, concepts and terms, generalizations, "
    "symbols and shapes, and skills. Write in English. Return JSON only."
)
SCHEMA_ANALYSIS = '''{
  "outcomes": ["Explain the meaning of...", "Describe the features of...", "Identify the main..."],
  "concepts": ["vocabulary item 1", "vocabulary item 2"],
  "generalizations": ["rule or generalization 1", "rule or generalization 2"],
  "symbols_shapes": [],
  "skills": ["Reads a short text and answers questions", "Writes a short paragraph"]
}'''


def semester_fallback(title):
    return {
        "outcomes": [f"By the end of the unit the student acquires the basic skills of {title}.",
                     "Apply the learning outcomes and the activities in the textbook."],
        "resources": ["Student's book", "Computers", "Data show"],
        "strategies": ["Direct instruction", "Group work", "Critical thinking"],
        "eval_strategies": ["Direct observation", "Performance-based assessment"],
        "eval_tools": ["Rating scale", "Worksheets"],
        "activities": ["Collaborative worksheets", "Presentations"],
    }


# ---------------------------------------------------------------- أدوات Word LTR
def _no_bidi(section):
    bidi = section._sectPr.find(qn("w:bidi"))
    if bidi is not None:
        section._sectPr.remove(bidi)


def _para(p, text, size=11, bold=False, align=L):
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    text = str(text)
    parts, pos = [], 0
    for m in re.finditer(r"\[\[(.+?)\]\]", text, flags=re.S):      # [[x]] = كلمة مسطَّرة في الأسئلة الوزارية
        if m.start() > pos:
            parts.append((text[pos:m.start()], False))
        parts.append((m.group(1), True))
        pos = m.end()
    if pos < len(text) or not parts:
        parts.append((text[pos:], False))
    r = None
    for seg, underline in parts:
        r = p.add_run(seg)
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.underline = underline
        r_pr = r._element.get_or_add_rPr()
        fonts = r_pr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            r_pr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(a), FONT)
    return r


def _borders(table, val="single"):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        if val == "single":
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "000000")
        b.append(el)
    tbl_pr.append(b)


def _table(doc, rows, cols, widths, borders=False):
    """جدول LTR طبيعي: الخلية 0 أقصى اليسار. عرض الأعمدة ثابت."""
    t = doc.add_table(rows=rows, cols=cols)
    t.autofit = False
    _borders(t, "single" if borders else "nil")
    tbl_pr = t._tbl.tblPr
    for tag, attrs in (("w:tblW", {"w:w": str(sum(int(w.twips) for w in widths)), "w:type": "dxa"}),
                       ("w:tblLayout", {"w:type": "fixed"}), ("w:jc", {"w:val": "left"})):
        el = tbl_pr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tbl_pr.append(el)
        for k, v in attrs.items():
            el.set(qn(k), v)
    grid = t._tbl.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.twips)))
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w
    return t


def _row_height(row, height):
    tr_pr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(height.twips)))
    h.set(qn("w:hRule"), "atLeast")
    tr_pr.append(h)


def _cell_borders(cell, visible=True):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcb = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom", "left", "right"):
        n = OxmlElement(f"w:{edge}")
        n.set(qn("w:val"), "single" if visible else "none")
        if visible:
            n.set(qn("w:sz"), "4")
            n.set(qn("w:space"), "0")
            n.set(qn("w:color"), "auto")
        tcb.append(n)
    tc_pr.append(tcb)


def _v_center(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), "center")
    tc_pr.append(v)


def _landscape(doc, first, top, bottom, side):
    section = doc.sections[0] if first else doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    section.top_margin, section.bottom_margin = Cm(top), Cm(bottom)
    section.left_margin = section.right_margin = Cm(side)
    _no_bidi(section)
    return section


def _footer(section, text, size=8):
    fp = section.footer.paragraphs[0]
    for r in list(fp.runs):
        r._element.getparent().remove(r._element)
    _para(fp, text, size, False, L)


def _doc_bytes(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _cell(cell, text, size=11, bold=False, align=L):
    return _para(cell.paragraphs[0], text, size, bold, align)


def _bullets(cell, items, size, first_is_para=True):
    for i, item in enumerate(items):
        p = cell.paragraphs[0] if (i == 0 and first_is_para) else cell.add_paragraph()
        _para(p, f"- {item}", size, False, L)
        p.paragraph_format.space_after = Pt(1)


# ---------------------------------------------------------------- خطة الدرس
STAGES = ["1- Engagement & Exploration", "2- Explanation & Interpretation", "3- Expansion & Support", "4- Evaluation"]


def _lesson_page(doc, first, header, plan):
    section = _landscape(doc, first, 0.8, 0.8, 2.35)
    _footer(section, "FORM #QF71-1-47 rev.b", 9)
    _para(doc.add_paragraph(), "Lesson Plan", 15, True, C)

    top = _table(doc, 1, 5, [Cm(6.0), Cm(5.5), Cm(6.0), Cm(5.0), Cm(2.5)])
    items = [f"Subject: {header['subject']}", f"Unit Title: {header['unit']}", f"Lesson Topic: {header['lesson']}",
             f"Prior Learning: {plan['prior_learning']}", f"Periods: {header['periods']}"]
    for cell, text, bold in zip(top.rows[0].cells, items, (True, True, True, False, False)):
        _cell(cell, text, 12, bold, L)

    out_t = _table(doc, 2, 1, [Cm(25.0)], borders=True)
    _row_height(out_t.rows[0], Cm(0.5))
    _row_height(out_t.rows[1], Cm(1.3))
    _cell(out_t.rows[0].cells[0], "Learning Outcomes:", 12.5, True)
    _bullets(out_t.rows[1].cells[0], plan["outcomes"], 12)

    main = _table(doc, 5, 4, [Cm(3.6), Cm(9.9), Cm(9.9), Cm(1.6)], borders=True)
    for cell, val in zip(main.rows[0].cells, ["Stage", "Teacher's Role", "Student's Role", "Time"]):
        _cell(cell, val, 12, True, C)
    for i, (label, stage) in enumerate(zip(STAGES, plan["stages"]), start=1):
        row = main.rows[i]
        _row_height(row, Cm(1.1))
        cells = row.cells
        _cell(cells[0], label, 10.5, True, C)
        _v_center(cells[0])
        _cell(cells[1], stage["teacher_role"], 12)
        _cell(cells[2], stage["student_role"], 12)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = sp.paragraph_format.space_after = Pt(2)

    widths = [Cm(14.0), Cm(0.5), Cm(3.5)] + [Cm(1.0)] * 7       # التأمل | فاصل | العناوين | 7 خلايا
    grid = doc.add_table(rows=4, cols=10)
    grid.autofit = False
    gp = grid._tbl.tblPr
    for tag, attrs in (("w:tblLayout", {"w:type": "fixed"}), ("w:jc", {"w:val": "left"})):
        el = gp.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            gp.append(el)
        for k, v in attrs.items():
            el.set(qn(k), v)
    for gc, w in zip(grid._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.twips)))
    for row in grid.rows:
        _row_height(row, Cm(0.65))
        for ci, (cell, w) in enumerate(zip(row.cells, widths)):
            cell.width = w
            _cell_borders(cell, visible=(ci != 1))
    for i, label in enumerate(["Day & Date", "Period No.", "Absent / Total", "Class / Section"]):
        _cell(grid.rows[i].cells[2], label, 11, True, C)
    reflect = grid.rows[0].cells[0]
    for r_idx in range(1, 4):
        reflect = reflect.merge(grid.rows[r_idx].cells[0])
    _cell_borders(reflect, True)
    _cell(reflect, "*Self-reflection on the teaching and learning processes:", 12, True)

    for _ in range(2):
        e = doc.add_paragraph()
        e.paragraph_format.space_before, e.paragraph_format.space_after = Pt(8), Pt(0)
    sig = _table(doc, 1, 4, [Cm(8.0), Cm(6.0), Cm(5.5), Cm(5.5)])
    vals = [f"Teacher (name & signature): {header['teacher']}".rstrip(), "Subject Specialist:", "Principal:",
            "School Development Advisor:"]
    for cell, val in zip(sig.rows[0].cells, vals):
        _cell(cell, val, 12)


def build_lesson_plan_docx(header, plans):
    doc = Document()
    for i, (lesson_header, plan) in enumerate(plans):
        _lesson_page(doc, i == 0, {**header, **lesson_header}, plan)
    return _doc_bytes(doc)


# ---------------------------------------------------------------- الخطة الفصلية
def _semester_page(doc, first, meta, unit):
    section = _landscape(doc, first, 0.6, 0.6, 1.0)
    _footer(section, "FORM # QF71-1-47rev.a")
    _para(doc.add_paragraph(), "Semester Plan", 13, True, C)

    top = _table(doc, 3, 3, [Cm(9.0), Cm(9.5), Cm(9.0)])
    r0, r1, r2 = (top.rows[i].cells for i in range(3))
    _cell(r0[0], f"Subject: {meta['subject']}", 10, True, L)
    _cell(r0[1], f"Semester: {meta['semester']}", 10, True, C)
    _cell(r0[2], f"Grade: {meta['grade']}", 10, True, R)
    _cell(r1[0], f"Unit Title: {unit['title']}", 10, True, L)
    _cell(r1[1], f"Pages: {unit.get('pages', '')}", 10, False, C)
    _cell(r1[2], f"Periods: {unit.get('periods', '')}", 10, False, R)
    _cell(r2[1], f"Duration: from {unit.get('start_date', '')} to {unit.get('end_date', '')}", 10, False, C)

    t = _table(doc, 3, 7, [Cm(7.2), Cm(4.5), Cm(3.0), Cm(3.0), Cm(3.2), Cm(3.5), Cm(3.3)], borders=True)
    ev = t.rows[0].cells[3].merge(t.rows[0].cells[4])
    _cell(ev, "Evaluation", 9.5, True, C)
    for col in (0, 1, 2, 5, 6):
        t.rows[0].cells[col].merge(t.rows[1].cells[col])
    heads = {0: "General Outcomes", 1: "Resources & Materials", 2: "Teaching Strategies",
             5: "Accompanying Activities", 6: "Self-reflection on the Unit"}
    for col, h in heads.items():
        _cell(t.rows[0].cells[col], h, 9.5, True, C)
    _cell(t.rows[1].cells[3], "Strategies", 9, True, C)
    _cell(t.rows[1].cells[4], "Tools", 9, True, C)
    for cell in t.rows[0].cells + t.rows[1].cells:
        _v_center(cell)

    plan = unit["plan"]
    body = t.rows[2].cells
    _cell(body[0], "By the end of the unit, the student should be able to:", 9, True)
    for item in plan.get("outcomes", []):
        _para(body[0].add_paragraph(), f"- {item}", 8.5)
    for col, key in ((1, "resources"), (2, "strategies"), (3, "eval_strategies"), (4, "eval_tools"), (5, "activities")):
        _bullets(body[col], plan.get(key, []), 8.5)
    _cell(body[6], "I am satisfied with: ..........................................\n\n"
                   "Challenges I faced: ..........................................\n\n"
                   "Suggestions for improvement: ..........................................", 8.5)

    info = doc.add_paragraph()
    info.paragraph_format.space_before, info.paragraph_format.space_after = Pt(6), Pt(4)
    _para(info, "General information about the students: ............................................................",
          9.5, True)
    sig = _table(doc, 2, 4, [Cm(8.0), Cm(7.0), Cm(6.0), Cm(6.0)])
    a, b = sig.rows[0].cells, sig.rows[1].cells
    _cell(a[0], f"Prepared by (teacher): {meta['teacher']}", 9.5, True)
    _cell(a[1], f"School principal: {meta['principal']}", 9.5, True)
    _cell(a[2], "Signature: ....................", 9)
    _cell(a[3], "Date: ...............", 9)
    _cell(b[0], "Educational supervisor: ....................", 9.5, True)
    _cell(b[2], "Signature: ....................", 9)
    _cell(b[3], "Date: ...............", 9)


def build_semester_docx(meta, units_plans):
    doc = Document()
    for i, u in enumerate(units_plans):
        _semester_page(doc, i == 0, meta, u)
    return _doc_bytes(doc)


# ---------------------------------------------------------------- تحليل المحتوى
def _analysis_page(doc, first, meta, unit):
    section = _landscape(doc, first, 0.8, 0.8, 2.0)
    _footer(section, "FORM # QF71-1-47rev")
    _para(doc.add_paragraph(), f"Content Analysis - {meta['semester']} Semester Plan - {meta['subject']}", 15, True, C)
    top = _table(doc, 2, 2, [Cm(12.5), Cm(12.5)])
    _cell(top.rows[0].cells[0], f"Grade: {meta['grade']}", 13, True, L)
    _cell(top.rows[0].cells[1], f"Academic year: {meta['year']}", 13, True, R)
    _cell(top.rows[1].cells[0], f"Unit Title: {unit['title']}", 13, True, L)
    _cell(top.rows[1].cells[1], f"Pages: {unit.get('pages', '')}", 13, True, R)

    t = _table(doc, 2, 6, [Cm(3.0), Cm(5.0), Cm(5.0), Cm(5.0), Cm(3.5), Cm(3.5)], borders=True)
    for cell, h in zip(t.rows[0].cells, ["Unit", "Outcomes", "Concepts & Terms", "Generalizations",
                                          "Symbols & Shapes", "Skills"]):
        _cell(cell, h, 13, True, C)
    a, row = unit["analysis"], t.rows[1].cells
    _cell(row[0], unit["title"], 13, True, C)
    _cell(row[1], "By the end of the unit, the student should be able to:", 13, True)
    for item in a.get("outcomes", []):
        _para(row[1].add_paragraph(), f"- {item}", 13)
    for col, key, dash in ((2, "concepts", False), (3, "generalizations", True),
                           (4, "symbols_shapes", True), (5, "skills", True)):
        for item in a.get(key, []):
            _para(row[col].add_paragraph(), f"- {item}" if dash else f"{item}", 13)


def build_analysis_docx(meta, units_analysis):
    doc = Document()
    for i, u in enumerate(units_analysis):
        _analysis_page(doc, i == 0, meta, u)
    return _doc_bytes(doc)


# ---------------------------------------------------------------- الاختبار (إنجليزي LTR)
EXAM_KIND_EN = {"الأول": "First", "الثاني": "Second", "النهائي": "Final"}
SEM_EN = {"الأول": "First", "الثاني": "Second"}
ORD_EN = ["One", "Two", "Three", "Four"]
INSTR_EN = {
    "mcq": "Choose the correct answer for each of the following:",
    "tf": "Write (T) if the statement is true and (F) if it is false:",
    "fill": "Complete the following sentences with suitable words:",
    "short": "Answer the following questions:",
}
LABEL_EN = {"mcq": "Multiple Choice", "tf": "True or False", "fill": "Fill in the Blanks", "short": "Short Answer"}


def _correct_en(q):
    if q["type"] == "mcq":
        return q["options"][q["correct_index"]]
    if q["type"] == "tf":
        return "True" if q["correct"] else "False"
    return q["model_answer"]


def _exam_image(doc, q, load_image):
    iid = q.get("image_id")
    data = load_image(iid) if (iid and load_image) else None
    if data:
        p = doc.add_paragraph()
        p.alignment = C
        try:
            p.add_run().add_picture(io.BytesIO(data), width=Cm(8))
        except Exception:  # noqa: BLE001
            pass


def build_exam_docx(header, selected_by_type, marks, logo_bytes, load_image=None):
    """ورقة اختبار إنجليزية: A4 عمودي، اتجاه LTR، تعليمات الأسئلة والإجابات بالإنجليزية.
    header: liwa, school, kind, subject, year, semester_label, duration, date, teacher (قيم عربية للنوع والفصل تُترجَم هنا)."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.3)
    sec.left_margin = sec.right_margin = Cm(1.5)
    _no_bidi(sec)

    if logo_bytes:
        lp = doc.add_paragraph()
        lp.alignment = C
        try:
            lp.add_run().add_picture(io.BytesIO(logo_bytes), width=Cm(2.6))
        except Exception:  # noqa: BLE001
            pass
    _para(doc.add_paragraph(), f"Directorate of Education - {header['liwa']}" if header.get("liwa")
          else "Directorate of Education", 13, True, C)
    _para(doc.add_paragraph(), header.get("school") or "School name", 13, True, C)
    kind = EXAM_KIND_EN.get(header.get("kind"), header.get("kind") or "")
    _para(doc.add_paragraph(), f"{kind} Exam - {header['subject']}", 13, True, C)
    sem = SEM_EN.get(header.get("semester_label"), header.get("semester_label") or "")
    year = f"Academic year {header['year']}" if header.get("year") else "Academic year ....... / ......."
    _para(doc.add_paragraph(), f"{sem} Semester     {year}", 12, False, C)

    t = _table(doc, 1, 3, [Cm(6), Cm(6), Cm(6)])
    n_cell, d_cell, du_cell = t.rows[0].cells
    _cell(n_cell, "Student's name: " + "_" * 10, 12, False, L)
    _cell(d_cell, f"Date: {header['date']}", 12, False, C)
    _cell(du_cell, f"Duration: {header['duration']}", 12, False, R)
    doc.add_paragraph()

    summary, o = [], 0
    for qtype in ("mcq", "tf", "fill", "short"):
        items = selected_by_type.get(qtype) or []
        if not items:
            continue
        ordinal = ORD_EN[o] if o < len(ORD_EN) else str(o + 1)
        o += 1
        ht = _table(doc, 1, 2, [Cm(14), Cm(4)])
        _cell(ht.rows[0].cells[0], f"Question {ordinal}: {INSTR_EN[qtype]}", 13, True, L)
        if marks.get(qtype):
            _cell(ht.rows[0].cells[1], f"({marks[qtype]} marks)", 13, True, R)
        entries, printed = [], set()
        for n, q in enumerate(items, start=1):
            if q.get("passage"):      # النص الرئيسي مرة واحدة قبل أول سؤال فرعي عليه
                pk = q.get("passage_key") or q["passage"][:40]
                if pk not in printed:
                    printed.add(pk)
                    _para(doc.add_paragraph(), q["passage"], 11)
            if qtype == "tf":
                _para(doc.add_paragraph(), f"{n}. {q['question']}   (      )", 12)
            else:
                _para(doc.add_paragraph(), f"{n}. {q['question']}", 12)
            _exam_image(doc, q, load_image)
            if qtype == "mcq":
                for i, opt in enumerate(q["options"]):
                    _para(doc.add_paragraph(), f"{'ABCD'[i]}) {opt}", 12)
            elif qtype == "short":
                for _ in range(2):
                    _para(doc.add_paragraph(), "." * 80, 12)
            entries.append((n, q))
        summary.append((ordinal, LABEL_EN[qtype], entries))
        doc.add_paragraph()

    et = _table(doc, 1, 3, [Cm(6), Cm(6), Cm(6)])
    _cell(et.rows[0].cells[0], "End of questions", 12, True, L)
    _cell(et.rows[0].cells[1], "Good luck", 12, True, C)
    teacher = (header.get("teacher") or "").strip()
    _cell(et.rows[0].cells[2], "Teacher: " + (teacher or "_" * 14), 12, True, R)

    doc.add_page_break()
    _para(doc.add_paragraph(), "Answer Key (Teacher's reference)", 13, True, L)
    for ordinal, label, entries in summary:
        doc.add_paragraph()
        _para(doc.add_paragraph(), f"Question {ordinal} ({label}):", 12, True)
        for n, q in entries:
            page = f" (page {q.get('page_printed') or q['page']})" if q.get("page") else ""
            _para(doc.add_paragraph(), f"{n}. {_correct_en(q)}{page}", 12)
    return _doc_bytes(doc)
