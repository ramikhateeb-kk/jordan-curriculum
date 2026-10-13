"""
attendance.py
تبويب المعلم "📋 دفتر الحضور والغياب": توليد دفتر الحضور والغياب كاملاً بصيغة Excel،
عبر تعبئة بيانات ورقة الغلاف باستخدام الأسماء المعرّفة (Defined Names) لضمان الاستقرار التام.
"""

import calendar
import datetime
import io

import openpyxl
import pandas as pd
import streamlit as st
from openpyxl.styles import PatternFill
from openpyxl.cell.cell import MergedCell

import accounts as acc


TEMPLATE_PATH = "templates/نموذج_الغياب.xlsx"
SHEET_NAME = "1"

REG_HEADER_ROW = 68
REG_DATA_START = 69
REG_DATA_MAX_END = 124
MAX_STUDENTS = 54

BLOCK_STARTS = [125, 183, 241, 299, 357, 415, 473, 531, 589, 647, 705]
BLOCK_HEIGHT = 58

# (الاسم العربي للشهر، رقم الشهر الميلادي، 0=يتبع السنة الأولى / 1=يتبع السنة الثانية)
MONTHS = [
    ("آب", 8, 0),
    ("ايلول", 9, 0),
    ("تشرين الاول", 10, 0),
    ("تشرين الثاني", 11, 0),
    ("كانون الاول", 12, 0),
    ("كانون الثاني", 1, 1),
    ("شباط", 2, 1),
    ("آذار", 3, 1),
    ("نيسان", 4, 1),
    ("ايار", 5, 1),
    ("حزيران", 6, 1),
]

DAY_COLUMNS = [
    "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q",
    "R", "S", "T", "U", "V", "Y", "Z", "AA", "AB", "AC", "AD",
    "AE", "AF", "AG", "AH", "AI", "AJ", "AK", "AL", "AM"
]

AR_WEEKDAY = {
    0: "الاثنين",
    1: "الثلاثاء",
    2: "الاربعاء",
    3: "الخميس",
    4: "الجمعة",
    5: "السبت",
    6: "الاحد",
}

WEEKEND = {4, 5}

GRAY_FILL = PatternFill(
    fill_type="solid",
    start_color="D9D9D9",
    end_color="D9D9D9"
)


# أعمدة قسم القيد
REG_COLS = {
    "serial": "F",
    "first": "G",
    "father": "I",
    "grandfather": "K",
    "family": "M",
    "birthplace": "S",
    "nationality": "U",
    "religion": "Y",
    "repeated_grades": "AD",
    "guardian_name": "AH",
    "guardian_relation": "AJ",
    "guardian_job": "AL",
    "phone": "AN",
    "residence": "AO",
}

FIELDS = [
    "first",
    "father",
    "grandfather",
    "family",
    "birthplace",
    "nationality",
    "religion",
    "repeated_grades",
    "guardian_name",
    "guardian_relation",
    "guardian_job",
    "phone",
    "residence",
]

SUMMARY_DATA_START = 766


# ----------------------------------------------------------------
# أعمدة جدول إدخال البيانات
EDITOR_COLUMNS = {
    "residence": "مكان السكن",
    "phone": "رقم الهاتف",
    "guardian_job": "عمله",
    "guardian_relation": "درجة القرابة",
    "guardian_name": "اسم ولي الأمر",
    "repeated_grades": "الصفوف التي أعادها",
    "religion": "الديانة",
    "nationality": "الجنسية",
    "birthplace": "مكان الولادة",
    "family": "اسم العائلة",
    "grandfather": "اسم الجد",
    "father": "اسم الأب",
    "first": "الاسم الأول",
}


def empty_roster_df():
    return pd.DataFrame({
        col: [""] * MAX_STUDENTS
        for col in EDITOR_COLUMNS
    })


def build_roster_from_df(df):
    roster = []

    for _, row in df.iterrows():

        name_parts = [
            str(row.get(k) or "").strip()
            for k in (
                "first",
                "father",
                "grandfather",
                "family"
            )
        ]

        if not any(name_parts):
            continue

        student = {
            k: str(row.get(k) or "").strip()
            for k in EDITOR_COLUMNS
        }

        roster.append(student)

        if len(roster) >= MAX_STUDENTS:
            break

    if not roster:
        raise ValueError(
            "يرجى إدخال بيانات طالب واحد على الأقل "
            "(الاسم الأول واسم الأب كحد أدنى)."
        )

    return roster


# ----------------------------------------------------------------
# أدوات الكتابة الآمنة مع الخلايا المدمجة

def _get_writable_cell(ws, coordinate):
    """
    الحصول على خلية قابلة للكتابة.

    إذا كانت الخلية المطلوبة جزءاً من نطاق مدمج وليست
    الخلية العلوية اليسرى، يتم إرجاع الخلية العلوية اليسرى
    من نطاق الدمج.
    """

    cell = ws[coordinate]

    # الخلية عادية وقابلة للكتابة
    if not isinstance(cell, MergedCell):
        return cell

    # الخلية جزء من Merge
    for merged_range in ws.merged_cells.ranges:

        if coordinate in merged_range:

            return ws.cell(
                row=merged_range.min_row,
                column=merged_range.min_col
            )

    # احتياطياً
    return cell


def _set_cell_value(ws, coordinate, value):
    """
    كتابة قيمة في خلية بشكل آمن.

    إذا كانت coordinate تشير إلى MergedCell،
    يتم الكتابة في الخلية العلوية اليسرى لنطاق الدمج.
    """

    cell = _get_writable_cell(
        ws,
        coordinate
    )

    cell.value = value


# ----------------------------------------------------------------
# دالة التعبئة عبر الأسماء المعرّفة

def write_by_defined_name(wb, name, value):
    """
    تكتب القيمة داخل الخلية المرتبطة بالـ Defined Name.

    إذا كان الـ Defined Name يشير إلى نطاق مدمج،
    يتم استخدام الخلية العلوية اليسرى.
    """

    if name not in wb.defined_names:
        return False

    defined_name = wb.defined_names[name]

    for sheet_name, coord in defined_name.destinations:

        ws = wb[sheet_name]

        # إزالة $
        clean_coord = coord.replace("$", "")

        # إذا كان النطاق مثلاً:
        # AM34:AP34
        # نأخذ أول خلية
        clean_coord = clean_coord.split(":")[0]

        _set_cell_value(
            ws,
            clean_coord,
            value
        )

        return True

    return False


# ----------------------------------------------------------------
# منطق التواريخ والتقويم

def _to_date(val):

    if isinstance(val, str):
        return datetime.date.fromisoformat(val)

    if isinstance(val, datetime.datetime):
        return val.date()

    return val


def _semester_for_date(d, cal):

    sem1_start = _to_date(
        cal["sem1_start"]
    )

    sem1_end = _to_date(
        cal["sem1_end"]
    )

    sem2_start = _to_date(
        cal["sem2_start"]
    )

    sem2_end = _to_date(
        cal["sem2_end"]
    )

    if sem1_start <= d <= sem1_end:
        return "الاول"

    if sem2_start <= d <= sem2_end:
        return "الثاني"

    return None


def _block_semester_label(month_idx, cal):

    _, greg_month, year_slot = MONTHS[month_idx]

    year = cal["year1"] + year_slot

    mid = datetime.date(
        year,
        greg_month,
        15
    )

    return _semester_for_date(
        mid,
        cal
    )


def _active_blocks(cal):

    active = []

    for i, (_, greg_month, year_slot) in enumerate(MONTHS):

        year = cal["year1"] + year_slot

        days_in_month = calendar.monthrange(
            year,
            greg_month
        )[1]

        for day in range(
            1,
            days_in_month + 1
        ):

            d = datetime.date(
                year,
                greg_month,
                day
            )

            if _semester_for_date(
                d,
                cal
            ):

                active.append(i)
                break

    return active


# ----------------------------------------------------------------
# كتابة أسماء الطلاب

def _write_names(
    ws,
    start_row,
    roster,
    with_serial=True
):

    for i, student in enumerate(roster):

        row = start_row + i

        if with_serial:

            _set_cell_value(
                ws,
                f"A{row}",
                i + 1
            )

        _set_cell_value(
            ws,
            f"B{row}",
            student.get("first") or ""
        )

        _set_cell_value(
            ws,
            f"C{row}",
            student.get("father") or ""
        )

        _set_cell_value(
            ws,
            f"D{row}",
            student.get("grandfather") or ""
        )

        _set_cell_value(
            ws,
            f"E{row}",
            student.get("family") or ""
        )


# ----------------------------------------------------------------
# تظليل عمود

def _shade_column(
    ws,
    block_start,
    col
):

    for row in range(
        block_start + 2,
        block_start + 2 + 1 + MAX_STUDENTS
    ):

        cell = ws[
            f"{col}{row}"
        ]

        # نضع الـ fill فقط إذا كانت الخلية
        # تسمح بذلك
        if isinstance(cell, MergedCell):
            continue

        cell.fill = GRAY_FILL


# ----------------------------------------------------------------
# توليد دفتر الحضور

def generate_register(
    roster,
    class_name,
    section_name,
    year_academic_label,
    cal,
    holidays,
    teacher
):

    wb = openpyxl.load_workbook(
        TEMPLATE_PATH
    )

    ws = wb[SHEET_NAME]

    class_section = (
        f"{class_name} {section_name}"
        .strip()
    )

    # ------------------------------------------------------------
    # 1. تعبئة بيانات ورقة الغلاف
    # ------------------------------------------------------------

    cover_data = {

        "YEAR_LABEL":
            year_academic_label,

        "DIRECTORATE":
            teacher.get("directorate") or "",

        "SCHOOL_NAME":
            teacher.get("school_name") or "",

        "CLASS_NAME":
            class_name,

        "SECTION_NAME":
            section_name,

        "TEACHER_NAME":
            teacher.get("name") or "",
    }

    for defined_name, value in cover_data.items():

        write_by_defined_name(
            wb,
            defined_name,
            value
        )

    # ------------------------------------------------------------
    # 2. تعبئة قسم القيد
    # ------------------------------------------------------------

    for i, student in enumerate(roster):

        row = REG_DATA_START + i

        _set_cell_value(
            ws,
            f"{REG_COLS['serial']}{row}",
            i + 1
        )

        for field in FIELDS:

            _set_cell_value(
                ws,
                f"{REG_COLS[field]}{row}",
                student.get(field) or ""
            )

    # ------------------------------------------------------------
    # 3. تعبئة الكتل الشهرية
    # ------------------------------------------------------------

    active = _active_blocks(cal)

    for idx in active:

        name, greg_month, year_slot = MONTHS[idx]

        year = cal["year1"] + year_slot

        days_in_month = calendar.monthrange(
            year,
            greg_month
        )[1]

        block_start = BLOCK_STARTS[idx]

        semester = (
            _block_semester_label(
                idx,
                cal
            )
            or ""
        )

        # --------------------------------------------------------
        # معلومات رأس الشهر
        # --------------------------------------------------------

        _set_cell_value(
            ws,
            f"A{block_start}",
            f"الصف والشعبة: {class_section}"
        )

        _set_cell_value(
            ws,
            f"I{block_start}",
            f"الشهر: {name}"
        )

        _set_cell_value(
            ws,
            f"AF{block_start}",
            f"الفصل: {semester}"
        )

        _set_cell_value(
            ws,
            f"AN{block_start}",
            f"السنة الدراسية: {year_academic_label}"
        )

        # --------------------------------------------------------
        # أيام الشهر
        # --------------------------------------------------------

        for day_idx, col in enumerate(
            DAY_COLUMNS,
            start=1
        ):

            weekday_coord = (
                f"{col}{block_start + 2}"
            )

            if day_idx > days_in_month:

                _set_cell_value(
                    ws,
                    weekday_coord,
                    ""
                )

                _shade_column(
                    ws,
                    block_start,
                    col
                )

                continue

            d = datetime.date(
                year,
                greg_month,
                day_idx
            )

            _set_cell_value(
                ws,
                weekday_coord,
                AR_WEEKDAY[d.weekday()]
            )

            if (
                d.weekday() in WEEKEND
                or d in holidays
                or _semester_for_date(
                    d,
                    cal
                ) is None
            ):

                _shade_column(
                    ws,
                    block_start,
                    col
                )

        # --------------------------------------------------------
        # أسماء الطلاب
        # --------------------------------------------------------

        _write_names(
            ws,
            block_start + 3,
            roster
        )

    # ------------------------------------------------------------
    # 4. جدول الخلاصة السنوي
    # ------------------------------------------------------------

    _write_names(
        ws,
        SUMMARY_DATA_START,
        roster
    )

    # ------------------------------------------------------------
    # حفظ الملف
    # ------------------------------------------------------------

    out = io.BytesIO()

    wb.save(out)

    return (
        out.getvalue(),
        len(active)
    )


# ----------------------------------------------------------------
# واجهة Streamlit

def render_attendance_builder(user, ctx=None):

    st.subheader(
        "📋 دفتر الحضور والغياب"
    )

    cal = acc.get_academic_calendar()

    if not cal:

        st.error(
            "⚠️ لم يحدّد المدير العام التقويم الدراسي بعد "
            "(تواريخ الفصلين). يرجى التواصل مع الإدارة."
        )

        return

    year_label = (
        f"{cal['year1']}/{cal['year1'] + 1}"
    )
    hijri_label = (acc.year_labels(cal) or {}).get("hijri", "")

    st.caption(
        f"السنة الدراسية الحالية: {year_label}م"
        + (f" ({hijri_label}هـ)" if hijri_label else "") + " — "
        f"الفصل الأول: {cal['sem1_start']} إلى {cal['sem1_end']} | "
        f"الفصل الثاني: {cal['sem2_start']} إلى {cal['sem2_end']}"
    )

    c1, c2 = st.columns(2)

    # الصف/التخصص من الصفوف المعتمدة لهذا المعلم (لا يكتبها يدويًا)، والشعبة يكتبها هو
    import catalog
    try:
        _tracks = catalog.list_tracks()
    except Exception:  # noqa: BLE001
        _tracks = []
    _opts = acc.class_options(user["id"], _tracks)
    if _opts:
        class_name = c1.selectbox("الصف / التخصص", [o["label"] for o in _opts])
    else:
        class_name = c1.text_input("الصف", placeholder="مثال: التاسع")

    section_name = c2.text_input(
        "الشعبة",
        placeholder="مثال: أ"
    )

    st.markdown(
        "#### أدخل بيانات الطلاب"
    )

    if "att_roster_df" not in st.session_state:

        st.session_state[
            "att_roster_df"
        ] = empty_roster_df()

    edited = st.data_editor(

        st.session_state[
            "att_roster_df"
        ],

        column_config={
            k: st.column_config.TextColumn(v)
            for k, v in EDITOR_COLUMNS.items()
        },

        column_order=list(
            EDITOR_COLUMNS
        ),

        num_rows="fixed",

        hide_index=True,

        key="att_editor",

        width="stretch",
    )

    if st.button(
        "🚀 إنشاء دفتر الحضور والغياب الكامل",
        type="primary"
    ):

        if (
            not class_name.strip()
            or not section_name.strip()
        ):

            st.error(
                "يرجى إدخال الصف والشعبة."
            )

        else:

            try:

                roster = build_roster_from_df(
                    edited
                )

                holidays = {
                    _to_date(
                        h["holiday_date"]
                    )
                    for h in acc.list_holidays()
                }

                data, n_blocks = generate_register(

                    roster,

                    class_name.strip(),

                    section_name.strip(),

                    year_label,

                    cal,

                    holidays,

                    user
                )

                st.session_state[
                    "att_result"
                ] = {

                    "data": data,

                    "n": len(roster),

                    "blocks": n_blocks,

                    "fname":
                        f"دفتر_الحضور_"
                        f"{class_name.strip()}_"
                        f"{section_name.strip()}_"
                        f"{year_label.replace('/', '-')}.xlsx"
                }

                st.success(
                    f"✅ تم إنشاء الدفتر لعدد "
                    f"{len(roster)} طالبًا بنجاح "
                    f"وتعبئة ورقة الغلاف."
                )

            except ValueError as e:

                st.session_state.pop(
                    "att_result",
                    None
                )

                st.error(
                    str(e)
                )

            except Exception as e:

                st.session_state.pop(
                    "att_result",
                    None
                )

                st.error(
                    f"حدث خطأ أثناء إنشاء الملف: {e}"
                )

    result = st.session_state.get(
        "att_result"
    )

    if result:

        st.download_button(

            "📥 تنزيل دفتر الحضور والغياب الكامل",

            result["data"],

            file_name=result["fname"],

            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),

            type="primary"
        )
