import io
import re
import streamlit as st
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from gtts import gTTS

from config import GEMINI_MODEL
from rag import get_client, select_context

# ⚠️ تأكد أن اسم العميل هنا يتطابق تماماً مع اسم الملف ومتغير Supabase لديك في المشروع
try:
    from db import supabase
except ImportError:
    try:
        from database import supabase
    except ImportError:
        supabase = None


# ---------------------------------------------------------------- التعامل مع Supabase
def save_guide_to_supabase(book_slug, lesson_title, subject, content):
    """حفظ أو تحديث دليل الشرح في قاعدة بيانات Supabase"""
    if not supabase:
        st.error("❌ لم يتم العثور على إعدادات اتصال Supabase (تأكد من ملف db.py).")
        return False

    try:
        data = {
            "book_slug": book_slug,
            "lesson_title": lesson_title,
            "subject": subject,
            "content": content
        }
        # إجراء Upsert بالتطابق مع القيد المزدوج (book_slug, lesson_title)
        res = supabase.table("teacher_guides").upsert(
            data, 
            on_conflict="book_slug, lesson_title"
        ).execute()
        return True
    except Exception as e:
        st.error(f"❌ خطأ أثناء الحفظ في Supabase: {e}")
        return False


def get_guide_from_supabase(book_slug, lesson_title):
    """استرجاع دليل الشرح المخزن مسبقاً من Supabase"""
    if not supabase:
        return None

    try:
        response = supabase.table("teacher_guides") \
            .select("content") \
            .eq("book_slug", book_slug) \
            .eq("lesson_title", lesson_title) \
            .execute()
        
        if response.data and len(response.data) > 0:
            return response.data[0]["content"]
    except Exception as e:
        st.warning(f"⚠️ تعذر قراءة البيانات من Supabase: {e}")
    return None


# ---------------------------------------------------------------- معالجة النصوص والتصدير
def clean_markdown_for_speech(text):
    text = re.sub(r'\*+|\#+|`+|_+', '', text)
    text = re.sub(r'\[.*?\]\(.*?\)', '', text)
    text = re.sub(r'\n+', ' ', text)
    return text.strip()


def text_to_speech(text, lang='ar'):
    try:
        clean_text = clean_markdown_for_speech(text)
        if not clean_text:
            return None
        clean_text = clean_text[:1000]
        tts = gTTS(text=clean_text, lang=lang, slow=False)
        fp = io.BytesIO()
        tts.write_to_fp(fp)
        fp.seek(0)
        return fp
    except Exception as e:
        st.error(f"خطأ في توليد الصوت: {e}")
        return None


def clean_output_text(text):
    """تنظيف شامل للنصوص والرموز المشفرة"""
    text = re.sub(r'&lt;\s*br\s*/?&gt;', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'<\s*br\s*/?>', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'</?[^>]+>', '', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    lines = [line.strip() for line in text.split('\n')]
    return '\n'.join(lines)


def set_rtl(paragraph_or_table):
    pPr = paragraph_or_table._element.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)


def export_to_word(title, content_text):
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    h = doc.add_heading(title, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(h)

    lines = content_text.split('\n')
    in_table = False
    table_lines = []

    for line in lines:
        if line.startswith('|'):
            in_table = True
            table_lines.append(line)
        else:
            if in_table:
                create_word_table(doc, table_lines)
                in_table = False
                table_lines = []

            if line.strip():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_rtl(p)
                clean_line = line.replace('**', '').replace('###', '').replace('##', '')
                if line.startswith('#'):
                    run = p.add_run(clean_line)
                    run.bold = True
                    run.font.size = Pt(13)
                elif line.startswith('- ') or line.startswith('* '):
                    p.add_run('• ' + clean_line.lstrip('-* '))
                else:
                    p.add_run(clean_line)

    if in_table:
        create_word_table(doc, table_lines)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def create_word_table(doc, lines):
    rows = [l.strip().strip('|').split('|') for l in lines if '---' not in l]
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT

    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            if c_idx < len(table.columns):
                cell = table.cell(r_idx, c_idx)
                cell.text = val.strip()
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_rtl(p)


# ---------------------------------------------------------------- الواجهة الرئيسية
def render_teach_help_builder(idx, slug, subject, flat_lessons):
    """واجهة مساعد الشرح (دليل المعلم الشامل)"""
    st.subheader("🧑‍🏫 مساعد الشرح (دليل المعلم الشامل)")
    st.caption("دليل حصة متكامل يراعي مستويات الطلاب بأسلوب موجز ومبسط، مع نماذج جداول واستراتيجيات تقويم وتدريس.")

    # حقن CSS شامل وقوي لإجبار المحاذاة يمين كاملة لكل عناصر Streamlit
    st.markdown("""
        <style>
        /* إجبار الاتجاه والمحاذاة من اليمين لكافة أجزاء التطبيق */
        div[data-testid="stMarkdownContainer"] {
            direction: rtl !important;
            text-align: right !important;
        }
        div[data-testid="stMarkdownContainer"] p, 
        div[data-testid="stMarkdownContainer"] h1, 
        div[data-testid="stMarkdownContainer"] h2, 
        div[data-testid="stMarkdownContainer"] h3, 
        div[data-testid="stMarkdownContainer"] h4,
        div[data-testid="stMarkdownContainer"] li {
            direction: rtl !important;
            text-align: right !important;
        }
        div[data-testid="stMarkdownContainer"] ul {
            direction: rtl !important;
            text-align: right !important;
            padding-right: 1.5rem !important;
            padding-left: 0rem !important;
            list-style-position: inside !important;
        }
        div[data-testid="stMarkdownContainer"] table {
            direction: rtl !important;
            margin-left: auto !important;
            margin-right: 0 !important;
            width: 100% !important;
        }
        div[data-testid="stMarkdownContainer"] th, 
        div[data-testid="stMarkdownContainer"] td {
            text-align: right !important;
            direction: rtl !important;
        }
        </style>
    """, unsafe_allow_html=True)

    if not flat_lessons:
        st.warning("لا توجد دروس معرفة في هذا الكتاب.")
        return

    selected_lesson_label = st.selectbox(
        "اختر الدرس لإعداد دليل مساعد الشرح له:",
        [l["label"] for l in flat_lessons],
        key="teach_help_lesson_select"
    )
    chosen_lesson = next(l for l in flat_lessons if l["label"] == selected_lesson_label)

    st.markdown("""
    **استراتيجيات التدريس المعتمدة:** التدريس المباشر | حل المشكلات والاستقصاء | التعلم الجماعي | التعلم بالنشاط والتفكير الناقد.  
    **استراتيجيات وأدوات التقويم:** التقويم على الأداء | القلم والورقة | الملاحظة | قائمة الرصد | سلم التقدير اللفظي | سجل وصف سير التعلم.
    """)

    # 1. محاولة قراءة الدليل المحفوظ سابقاً من Supabase
    saved_guide = get_guide_from_supabase(slug, chosen_lesson["title"])

    if st.button("🚀 إنشاء دليل مساعد الشرح", type="primary", key="generate_teach_help_btn"):
        with st.spinner("جارٍ إعداد خطة ودليل الشرح التفصيلي للحصة..."):
            try:
                pages = (chosen_lesson["start"], chosen_lesson["end"])
                ctx_pages = select_context(idx, chosen_lesson["title"], pages)
                context_text = "\n".join([f"صفحة {p['page']}: {p['text']}" for p in ctx_pages])

                prompt = f"""
أنت مشرف تربوي وخبير تدريس للمنهاج الأردني.
قم بإنشاء دليل شرح مبسط وموجز للمعلم لدرس: "{chosen_lesson['title']}" في مادة {subject}.

نص الدرس من الكتاب:
{context_text}

شروط الصياغة والتنسيق الهامة جداً (التزام صارم):
1. ممنوع منعاً باتاً استخدام أية وسوم HTML مثل <br> أو <b> في النص أو داخل خلايا الجداول نهائياً.
2. للفصل بين النقاط داخل الخلية الواحدة في الجدول، استخدم علامة الشرطة (-) فقط.
3. كن موجزاً ومباشراً وركز على النقاط العملية والخطوات التنفيذية فقط (بدون إطالة أو حشو).
4. صمم الجداول بشكل مكتمل دون إهمال أو ترك خانات فارغة.

المحاور المطلوبة:

1. **جدول توزيع زمن الحصة (45 دقيقة):**
جدول أفق يضم: (المرحلة | الزمن | الأنشطة المقترحة | الموارد والأدوات).
المراحل: التهيئة والإثارة (5د) ، العرض والتوضيح (15د) ، الأنشطة الموجهة والجماعية (15د) ، التقويم والختام (10د).

2. **خطوات شرح الدرس والتمايز (موجز ومباشر):**
- **للطلبة ضعاف التحصيل:** (أمثلة حسية، تدرج، تبسيط مفاهيم).
- **للطلبة المجتهدين:** (تحدي استقصائي، سؤال تفكير ناقد).

3. **تطبيق استراتيجيات التدريس الحديثة:**
نقاط مكثفة لتطبيق: التدريس المباشر، حل المشكلات والتعلم الجماعي.

4. **نماذج أدوات التقويم المتبعة في الدرس:**
قم بتقديم **نموذج جدول مكتمل جاهز للاستخدام** لإحدى الأدوات التالية المتناسبة مع الدرس:
- **نموذج قائمة رصد (Checklist)** (المعيار | نعم | لا | ملاحظات)
أو 
- **نموذج سلم تقدير لفظي (Rubric)** (المعيار | ممتاز | جيد | بحاجة لدعم)
"""
                client = get_client()
                response = client.models.generate_content(
                    model=GEMINI_MODEL,
                    contents=prompt
                )

                result_text = clean_output_text(response.text)

                # 2. حفظ النتيجة في Supabase
                success = save_guide_to_supabase(slug, chosen_lesson["title"], subject, result_text)
                saved_guide = result_text
                
                if success:
                    st.success("✅ تم توليد الدليل وحفظه بنجاح في قاعدة البيانات!")

            except Exception as e:
                st.error(f"حدث خطأ أثناء بناء دليل الشرح: {e}")

    # 3. عرض النتائج والتنزيل
    if saved_guide:
        st.markdown("### 📋 دليل مساعد الشرح للحصة:")
        st.markdown(saved_guide)

        st.markdown("---")
        col_download, col_audio = st.columns(2)

        with col_download:
            doc_file = export_to_word(f"دليل شرح: {chosen_lesson['title']}", saved_guide)
            st.download_button(
                label="📥 تحميل الدليل كملف Word (.docx)",
                data=doc_file,
                file_name=f"دليل_شرح_{chosen_lesson['title']}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

        with col_audio:
            if st.checkbox("🔊 الاستماع لملخص الدليل"):
                with st.spinner("جارٍ توليد الصوت..."):
                    audio_fp = text_to_speech(saved_guide, lang='ar')
                    if audio_fp:
                        st.audio(audio_fp, format='audio/mp3')
