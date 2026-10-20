import hashlib
import io
import json
import re
import pptx
import requests
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
import streamlit as st

import catalog
from config import GEMINI_MODEL
from rag import get_client, select_context
import auth_ui
from file_store import upload_file

TABLE = "lesson_slides"
TIMEOUT = 30
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

# ألوان وتصميم العرض التقديمي الجذاب
DARK_NAVY = RGBColor(30, 58, 95)     # #1e3a5f
TEAL = RGBColor(46, 111, 94)         # #2e6f5e
GOLD = RGBColor(212, 160, 23)        # #d4a017
LIGHT_BG = RGBColor(248, 249, 250)   # #f8f9fa
WHITE = RGBColor(255, 255, 255)


FILE_VERSION = "v2-rtl"   # غيّره عند تعديل تصميم الملف ليُعاد بناء الملفات المحفوظة تلقائيًا
ARABIC_FONT = "Arial"


def _add_text(paragraph, text, size, color, bold=False, align=PP_ALIGN.RIGHT, bullet=None):
    """يضيف نصًا عربيًا في فقرة بتنسيق RTL كامل (اتجاه الفقرة + محاذاة + خط للغات المعقدة).
    bullet: رمز النقطة (يوضع في Run مستقل حتى لا يتأثر تشكيل الحروف العربية)."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("rtl", "1")
    paragraph.alignment = align

    if bullet:
        b = paragraph.add_run()
        b.text = bullet + " "
        b.font.size = Pt(size)
        b.font.color.rgb = color

    run = paragraph.add_run()
    run.text = text
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.name = ARABIC_FONT
    f.color.rgb = color

    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "ar-SA")
    cs = rPr.find(qn("a:cs"))
    if cs is None:
        cs = rPr.makeelement(qn("a:cs"), {})
        rPr.append(cs)
    cs.set("typeface", ARABIC_FONT)
    return run


def create_pptx_presentation(slide_data, topic_title):
    """إنشاء عرض تقديمي عربي (من اليمين إلى اليسار) باستخدام python-pptx"""
    prs = Presentation()
    prs.slide_width = Inches(13.33)  # Wide Screen 16:9
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # 1. صفحة العنوان الرئيسي
    title_slide = prs.slides.add_slide(blank_layout)
    fill = title_slide.background.fill
    fill.solid()
    fill.fore_color.rgb = DARK_NAVY

    tx_box = title_slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.33), Inches(3))
    tf = tx_box.text_frame
    tf.word_wrap = True

    _add_text(tf.paragraphs[0], "📚 عرض تقديمي مبسط", 28, GOLD, bold=True, align=PP_ALIGN.CENTER)
    _add_text(tf.add_paragraph(), topic_title, 44, WHITE, bold=True, align=PP_ALIGN.CENTER)
    _add_text(tf.add_paragraph(), "مساعد المنهاج الأردني الذكي", 20, RGBColor(220, 225, 230),
              align=PP_ALIGN.CENTER)

    # 2. شرائح المحتوى الشارح المبسط
    for slide_info in slide_data:
        slide = prs.slides.add_slide(blank_layout)
        bg_fill = slide.background.fill
        bg_fill.solid()
        bg_fill.fore_color.rgb = LIGHT_BG

        # العنوان (محاذاة يمين)
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.73), Inches(1))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        _add_text(tf_title.paragraphs[0], slide_info.get("title", ""), 30, TEAL, bold=True)

        # المحتوى (محاذاة يمين، والنقطة 🔹 تظهر على اليمين)
        content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.73), Inches(5))
        tf_content = content_box.text_frame
        tf_content.word_wrap = True

        for i, pt in enumerate(slide_info.get("points", [])):
            p = tf_content.paragraphs[0] if i == 0 else tf_content.add_paragraph()
            _add_text(p, pt, 22, DARK_NAVY, bullet="🔹")
            p.space_after = Pt(14)

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------- التعامل مع Supabase (REST)
def _use_supabase():
    return catalog.backend_name() == "supabase"


def get_slides_row(slug, page_start, page_end):
    """يعيد {data, file_url, file_path} للعرض المحفوظ لهذا الدرس، أو None."""
    if not _use_supabase():
        return None
    try:
        url, h = catalog._sb()
        r = requests.get(f"{url}/rest/v1/{TABLE}", headers=h, timeout=TIMEOUT, params={
            "select": "data,file_url,file_path", "slug": f"eq.{slug}",
            "page_start": f"eq.{page_start}", "page_end": f"eq.{page_end}"})
        catalog._check(r)
        rows = r.json()
        return rows[0] if rows else None
    except Exception as e:  # noqa: BLE001
        st.warning(f"⚠️ تعذر الاستعلام من قاعدة البيانات: {e}")
        return None


def save_slides_row(slug, page_start, page_end, lesson_title, subject, data,
                    file_url=None, file_path=None):
    """حفظ/تحديث (Upsert) محتوى الشرائح مع رابط الملف إن وُجد."""
    if not _use_supabase():
        return False
    try:
        row = {"slug": slug, "page_start": page_start, "page_end": page_end,
               "lesson_title": lesson_title, "subject": subject, "data": data}
        if file_url:
            row["file_url"] = file_url
            row["file_path"] = file_path
        url, h = catalog._sb()
        r = requests.post(f"{url}/rest/v1/{TABLE}",
                          headers={**h, "Prefer": "resolution=merge-duplicates,return=minimal"},
                          timeout=TIMEOUT, params={"on_conflict": "slug,page_start,page_end"},
                          data=json.dumps([row]))
        catalog._check(r)
        return True
    except Exception as e:  # noqa: BLE001
        st.error(f"❌ خطأ أثناء الحفظ في قاعدة بيانات Supabase: {e}")
        return False


def build_file_name(title):
    safe = re.sub(r'[\\/:*?"<>|]', "", f"عرض_تقديمي_{title}").replace(" ", "_")
    return safe + ".pptx"


def expected_file_path(slug, page_start, page_end, slide_data):
    """مسار الملف في Storage؛ يتغير عند تغيّر المحتوى أو FILE_VERSION (فيُعاد بناء الملف تلقائيًا)."""
    key = f"{slug}_{page_start}_{page_end}"
    raw = FILE_VERSION + json.dumps(slide_data, ensure_ascii=False, sort_keys=True)
    digest = hashlib.md5(raw.encode("utf-8")).hexdigest()[:8]
    return f"slides/{hashlib.md5(key.encode('utf-8')).hexdigest()[:16]}_{digest}.pptx"


def publish_slides(slug, page_start, page_end, lesson_title, subject, slide_data):
    """ينشئ ملف PPTX، يرفعه إلى Storage، ويحفظ الشرائح ورابط الملف في الجدول. يعيد الرابط أو None."""
    file_link, file_path = None, None
    if _use_supabase():
        try:
            pptx_bytes = create_pptx_presentation(slide_data, lesson_title).getvalue()
            file_path = expected_file_path(slug, page_start, page_end, slide_data)
            file_link = upload_file(file_path, pptx_bytes, PPTX_MIME, build_file_name(lesson_title))
        except Exception as e:  # noqa: BLE001
            st.warning(f"⚠️ تم حفظ الشرائح لكن تعذر إنشاء رابط التحميل: {e}")
    save_slides_row(slug, page_start, page_end, lesson_title, subject, slide_data, file_link, file_path)
    return file_link


def _generate_slide_data(idx, subject, chosen_lesson, user=None):
    pages = (chosen_lesson["start"], chosen_lesson["end"])
    ctx_pages = select_context(idx, chosen_lesson["title"], pages)
    context_text = "\n".join([f"صفحة {p['page']}: {p['text']}" for p in ctx_pages])

    prompt = f"""
أنت خبير إعداد المناهج والعروض التقديمية المبسطة.
قم بإنشاء هيكل عرض تقديمي متكامل لدرس: "{chosen_lesson['title']}" في مادة {subject}.
استعن بالنص التالي من الكتاب:
{context_text}

الشروط:
1. قم بتبسيط المفاهيم والشرح بأقصى درجة ممكنة ليكون موجهاً للطلاب ضعاف التحصيل.
2. يتكون العرض من (5 إلى 7) شرائح.
3. يجب أن تكون مخرجاتك بنظام JSON حصراً، على الشكل التالي:
[
  {{
    "title": "عنوان الشريحة الأول",
    "points": ["نقطة شارحة 1 بكلمات بسيطة", "نقطة شارحة 2 بأسلوب مباشر", "مثال توضيحي"]
  }}
]
أرجع كود JSON فقط بدون أي مقدمات أو علامات إضافية.
"""
    if user is not None:
        auth_ui.require_quota_or_stop(user)
    client = get_client()
    response = client.models.generate_content(model=GEMINI_MODEL, contents=prompt)

    raw_json = response.text.strip()
    if raw_json.startswith("```json"):
        raw_json = raw_json[7:-3].strip()
    elif raw_json.startswith("```"):
        raw_json = raw_json[3:-3].strip()
    return json.loads(raw_json)


# ---------------------------------------------------------------- الواجهة
def render_slides_builder(idx, slug, subject, flat_lessons, user=None):
    """واجهة المعلم: أول طلب = توليد + حفظ + رابط تحميل، والطلبات التالية = استرجاع من قاعدة البيانات."""
    st.subheader("💻 عرض تقديمي (PowerPoint - PPTX)")
    st.caption("عرض تقديمي بتصميم جذاب ونصوص مبسطة تناسب الطلبة ضعاف التحصيل وقابلة للتعديل الكامل.")

    if not flat_lessons:
        st.warning("لا توجد دروس معرفة في هذا الكتاب.")
        return

    selected_lesson_label = st.selectbox(
        "اختر الدرس لإعداد العرض التقديمي له:",
        [l["label"] for l in flat_lessons],
        key="slides_lesson_select"
    )
    chosen_lesson = next(l for l in flat_lessons if l["label"] == selected_lesson_label)
    title = chosen_lesson["title"]
    p_start, p_end = chosen_lesson["start"], chosen_lesson["end"]
    lesson_key = f"{slug}_{p_start}_{p_end}"

    # العرض الحالي (يبقى ظاهرًا بعد أي إعادة تشغيل للصفحة)
    shown = st.session_state.get("slides_shown", {})
    slide_data = shown.get("data") if shown.get("key") == lesson_key else None
    file_url = shown.get("file_url") if shown.get("key") == lesson_key else None

    if st.button("✨ عرض / توليد العرض التقديمي (PPTX)", type="primary", key="slides_generate_btn"):
        with st.spinner("جارٍ تحضير العرض التقديمي..."):
            # 1) أولاً: هل العرض محفوظ في قاعدة البيانات؟ → استرجاعه دون أي طلب إلى Gemini
            row = get_slides_row(slug, p_start, p_end)
            if row and row.get("data"):
                slide_data = row["data"]
                file_url = row.get("file_url")
                st.info("💡 تم استرجاع العرض من قاعدة البيانات (لم يُولَّد من جديد).")
                if _use_supabase() and row.get("file_path") != expected_file_path(slug, p_start, p_end, slide_data):
                    # الملف المحفوظ بتصميم قديم → إعادة بنائه ورفعه (بدون أي طلب إلى Gemini)
                    file_url = publish_slides(slug, p_start, p_end, title, subject, slide_data)
            else:
                # 2) أول طلب: توليد بـ Gemini ثم حفظ الشرائح وإنشاء ملف PPTX ورابطه
                try:
                    slide_data = _generate_slide_data(idx, subject, chosen_lesson, user=user)
                    file_url = publish_slides(slug, p_start, p_end, title, subject, slide_data)
                    st.success("🎉 تم إنشاء العرض التقديمي وحفظه في قاعدة البيانات!")
                except Exception as e:  # noqa: BLE001
                    slide_data = None
                    st.error(f"حدث خطأ أثناء بناء العرض التقديمي: {e}")

        if slide_data:
            st.session_state["slides_shown"] = {"key": lesson_key, "data": slide_data, "file_url": file_url}

    if slide_data:
        # عرض محفوظ بدون ملف/رابط → زر لإنشائه بدون طلب Gemini
        if not file_url and _use_supabase():
            if st.button("🔗 إنشاء ملف PPTX ورابط تحميل لهذا العرض", key="slides_make_link_btn"):
                with st.spinner("جارٍ إنشاء الملف ورفعه..."):
                    file_url = publish_slides(slug, p_start, p_end, title, subject, slide_data)
                    st.session_state["slides_shown"] = {"key": lesson_key, "data": slide_data, "file_url": file_url}

        st.markdown("### 🖼️ معاينة محتوى الشرائح:")
        for i, slide_info in enumerate(slide_data, start=1):
            with st.expander(f"الشريحة {i}: {slide_info.get('title', '')}", expanded=True):
                for pt in slide_info.get("points", []):
                    st.write(f"• {pt}")

        st.markdown("---")
        if file_url:
            st.link_button("🔗 تحميل العرض (رابط مباشر)", file_url, type="primary")
            st.caption("انسخ الرابط لمشاركته:")
            st.code(file_url, language=None)

        st.download_button(
            label="📥 تحميل العرض التقديمي (.pptx)",
            data=create_pptx_presentation(slide_data, title),
            file_name=build_file_name(title),
            mime=PPTX_MIME,
        )
