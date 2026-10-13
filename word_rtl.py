"""
أدوات مشتركة لإنتاج ملفات Word عربية من اليمين إلى اليسار (يستخدمها exam.py وlesson_plan.py وأي
مولّد Word آخر لاحقًا).

الطريقة المعتمدة (مُختبَرة وتعمل فعليًا، بخلاف محاولات سابقة اعتمدت على w:bidi وw:bidiVisual):
- الجدول نفسه يُترَك باتجاه LTR صراحة (w:tblDir="ltr")، تمامًا كما يفعل Word عمليًا عند الضغط على
  «اتجاه النص من اليسار إلى اليمين» يدويًا على جدول عربي. الخلية الأولى (index 0) تكون فعليًا أقصى
  اليسار، والأخيرة أقصى اليمين — هذا معكوس بديهيًا لكنه الترتيب الصحيح فعليًا.
- كل فقرة: لا نضيف لها w:bidi إطلاقًا (كان هذا مصدر تضارب سابقًا)، بل نضبط محاذاتها صراحة عبر w:jc
  (يمين/وسط/يسار) ونُصفِّر أي إزاحة موروثة (w:ind).
- كل "تشغيلة نص" (run): تُوسَم بـ Font.rtl=True وFont.complex_script=True مع خط عربي في rFonts
  (ascii وhAnsi وcs معًا)، ليتصرف Word مع رموزها (الأقواس، النقاط) بالشكل الصحيح.
- أي سطر نص عادي يجب أن يصل فعليًا لحافة الصفحة (يمين أو يسار) يُبنى كجدول بخلية واحدة بعرض منطقة
  الكتابة كاملة (right()/left())، لا كفقرة حرة في متن المستند — فالفقرات الحرة عرضة لإزاحة لا يمكن
  ضبطها بثقة، بينما الجدول بعرض صريح يضمن وصول النص لحافة الصفحة فعليًا.
"""
import io

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ARABIC_FONT = "Arial"


def get_or_add(parent, tag):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    return el


def rtl_section(section):
    """اتجاه القسم في Word من اليمين إلى اليسار."""
    if section._sectPr.find(qn("w:bidi")) is None:
        section._sectPr.append(OxmlElement("w:bidi"))


def rtl_table(table):
    """يُجبر اتجاه الجدول على LTR صراحة (لا bidiVisual). راجع شرح أعلى الملف."""
    tbl_pr = table._tbl.tblPr
    bidi_visual = tbl_pr.find(qn("w:bidiVisual"))
    if bidi_visual is not None:
        tbl_pr.remove(bidi_visual)
    tbl_dir = get_or_add(tbl_pr, "w:tblDir")
    tbl_dir.set(qn("w:val"), "ltr")


def no_borders(table):
    _set_borders(table, "nil")


def visible_borders(table):
    """حدود رفيعة سوداء ظاهرة لكل أضلاع الجدول وخطوطه الداخلية (لجداول الشبكة كجدول مراحل الدرس)."""
    _set_borders(table, "single", sz="4", color="000000")


def _set_borders(table, val, sz=None, color=None):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        if sz:
            el.set(qn("w:sz"), sz)
        if color:
            el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def rtl_p(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    """محاذاة صريحة (لا w:bidi)، مع تصفير أي إزاحة موروثة."""
    p_pr = p._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is not None:
        p_pr.remove(bidi)

    jc = get_or_add(p_pr, "w:jc")
    jc.set(qn("w:val"), {WD_ALIGN_PARAGRAPH.CENTER: "center", WD_ALIGN_PARAGRAPH.LEFT: "left"}
          .get(align, "right"))

    ind = get_or_add(p_pr, "w:ind")
    for attr in ("w:left", "w:right", "w:firstLine", "w:hanging"):
        ind.set(qn(attr), "0")

    pf = p.paragraph_format
    pf.left_indent = pf.right_indent = pf.first_line_indent = Cm(0)
    pf.space_before = pf.space_after = Pt(0)
    return p


def run(p, text, size=14, bold=False, underline=False):
    """إضافة نص عربي مع إجبار Word على التعامل معه كسياق RTL صحيح."""
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.underline = underline
    r.font.name = ARABIC_FONT
    r.font.rtl = True
    r.font.complex_script = True
    r_pr = r._element.get_or_add_rPr()
    r_fonts = get_or_add(r_pr, "w:rFonts")
    r_fonts.set(qn("w:ascii"), ARABIC_FONT)
    r_fonts.set(qn("w:hAnsi"), ARABIC_FONT)
    r_fonts.set(qn("w:cs"), ARABIC_FONT)
    return r


def center(doc, text, size=13, bold=True):
    p = doc.add_paragraph()
    rtl_p(p, WD_ALIGN_PARAGRAPH.CENTER)
    run(p, text, size, bold)
    return p


def doc_table(doc, rows, cols, col_widths, borders=False):
    """جدول داخل متن المستند، بعرض ومحاذاة يمين مضبوطَين صراحة. الجدول LTR (راجع rtl_table)، فالخلية
    index 0 تكون فعليًا أقصى اليسار، والأخيرة أقصى اليمين.
    borders=False (الافتراضي): حدود مخفية، لصفوف معلومات تُقرأ كسطر عادي.
    borders=True: حدود ظاهرة، لجداول الشبكة الفعلية (كجدول مراحل الدرس)."""
    t = doc.add_table(rows=rows, cols=cols)
    t.autofit = False
    rtl_table(t)
    visible_borders(t) if borders else no_borders(t)

    tbl_pr = t._tbl.tblPr
    total_twips = sum(int(w.twips) for w in col_widths)
    tbl_w = get_or_add(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")
    jc = get_or_add(tbl_pr, "w:jc")
    jc.set(qn("w:val"), "right")
    tbl_layout = get_or_add(tbl_pr, "w:tblLayout")
    tbl_layout.set(qn("w:type"), "fixed")

    grid = t._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for grid_col, w in zip(grid.findall(qn("w:gridCol")), col_widths):
            grid_col.set(qn("w:w"), str(int(w.twips)))
    for row_ in t.rows:
        for cell, w in zip(row_.cells, col_widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = get_or_add(tc_pr, "w:tcW")
            tc_w.set(qn("w:w"), str(int(w.twips)))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = w
    return t


def right(doc, text, width, size=12, bold=False):
    """سطر بعرض `width` كاملًا (جدول بخلية واحدة) ثم محاذاة النص إلى أقصى اليمين فعليًا."""
    t = doc_table(doc, 1, 1, [width])
    p = rtl_p(t.rows[0].cells[0].paragraphs[0], WD_ALIGN_PARAGRAPH.RIGHT)
    run(p, text, size, bold)
    return p


def left(doc, text, width, size=12, bold=False):
    """سطر بعرض `width` كاملًا ومحاذاة أقصى اليسار فعليًا."""
    t = doc_table(doc, 1, 1, [width])
    p = rtl_p(t.rows[0].cells[0].paragraphs[0], WD_ALIGN_PARAGRAPH.LEFT)
    run(p, text, size, bold)
    return p


def vertical_cell(cell, direction="btLr"):
    """يدوّر نص الخلية عموديًا (لعناوين أعمدة ضيقة طويلة، كما في نموذج خطة الدرس)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    d = OxmlElement("w:textDirection")
    d.set(qn("w:val"), direction)
    tc_pr.append(d)
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "center")
    tc_pr.append(va)


def add_picture_centered(doc, image_bytes, width=Cm(2.6)):
    """شعار أو صورة متوسطة أعلى الصفحة، بلا توقف إن كانت الصورة غير صالحة."""
    p = doc.add_paragraph()
    rtl_p(p, WD_ALIGN_PARAGRAPH.CENTER)
    try:
        p.add_run().add_picture(io.BytesIO(image_bytes), width=width)
    except Exception:  # noqa: BLE001
        pass
    return p
